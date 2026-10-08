import os, re, json, sqlite3, hmac, secrets, threading, time, base64
from datetime import datetime
from flask import Flask, render_template, request, jsonify, redirect, session, url_for, send_file, abort, g
from db import get_conn, add_timeline, DB_PATH, init_db, get_user_secrets
from ai_engine import generate_cover_letter, generate_linkedin_note, score_job, get_fallback_gemini_key
import voice_engine as ve
import quiz_coach
from notifier import send_email_digest, recipient_for_user
from auth import signup_user, login_user
from crypto_util import get_app_secret, encrypt_secret

app = Flask(__name__, template_folder='.')
app.secret_key = get_app_secret()
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    # Secure cookies by default; set SESSION_COOKIE_SECURE=0 for plain-http local dev
    SESSION_COOKIE_SECURE=os.getenv("SESSION_COOKIE_SECURE", "1") != "0",
    MAX_CONTENT_LENGTH=int(os.getenv("MAX_UPLOAD_MB", "10")) * 1024 * 1024,
)

# The "Dev Mode" fake Google/LinkedIn sign-in lets anyone log in as any email.
# It only works when explicitly enabled for local testing.
ALLOW_DEV_LOGIN = os.getenv("ALLOW_DEV_LOGIN", "0") == "1"

# How long onboarding waits for the first job sync before returning (keep well under the gunicorn timeout)
ONBOARD_WAIT_SECONDS = int(os.getenv("ONBOARD_WAIT_SECONDS", "25"))

# Shown in the settings form instead of a stored secret; posting it back means "unchanged"
SECRET_PLACEHOLDER = "********"

# Initialize database on startup (crucial for Gunicorn/Render deployments)
init_db(DB_PATH)

# Set session cookies lifetime to be long so login stays active
from datetime import timedelta
app.permanent_session_lifetime = timedelta(days=30)

def get_user_id() -> int:
    return session.get("user_id")

def is_admin() -> bool:
    return session.get("user_id") == 1

def get_user_settings(conn, user_id, mask_secrets=False):
    row = conn.execute("SELECT resume_text, imap_email, imap_password, gemini_api_key, linkedin_profile, name, designation, share_profile, resume_filename, weight_thumbs_up, weight_applied, weight_thumbs_down, weight_rejected FROM users WHERE id = ?", (user_id,)).fetchone()
    if not row:
        return {}
    settings = dict(row)
    if mask_secrets:
        for k in ("imap_password", "gemini_api_key"):
            settings[k] = SECRET_PLACEHOLDER if settings.get(k) else ""
    else:
        settings.update({k: v for k, v in get_user_secrets(conn, user_id).items() if k != "imap_email"})
    return settings

# Background jobs are de-duplicated per user so repeated clicks don't stack up LLM calls
_running_jobs = set()
_running_lock = threading.Lock()

def run_in_background(key, fn, *args):
    with _running_lock:
        if key in _running_jobs:
            return False
        _running_jobs.add(key)

    def runner():
        try:
            fn(*args)
        except Exception as e:
            print(f"Background job {key} failed: {e}")
        finally:
            with _running_lock:
                _running_jobs.discard(key)

    threading.Thread(target=runner, daemon=True).start()
    return True

def rescore_inbox_in_background(uid):
    def work():
        conn = get_conn(DB_PATH)
        conn.execute(
            "UPDATE jobs SET status = 'new' WHERE user_id = ? AND status IN ('new', 'scored', 'ready')",
            (uid,)
        )
        conn.commit()
        conn.close()
        from ai_engine import process_new_jobs
        process_new_jobs(DB_PATH, min_score=0, user_id=uid)
    return run_in_background(("rescore", uid), work)

def _check_oauth_state():
    expected = session.pop("oauth_state", None)
    got = request.args.get("state", "")
    return bool(expected) and hmac.compare_digest(expected, got)

@app.before_request
def require_login():
    allowed_endpoints = ["login", "signup", "static", "index", "auth_google", "auth_google_callback", "auth_linkedin", "auth_linkedin_callback", "auth_mock_callback", "serve_logo", "serve_favicon", "serve_hero_background", "cron_daily_recommendations"]
    if not session.get("user_id"):
        if request.endpoint and request.endpoint not in allowed_endpoints:
            if request.path.startswith("/api/"):
                return jsonify({"ok": False, "error": "Not signed in"}), 401
            return redirect(url_for("login"))
    # The login cookie is shared by every tab. A board page opened as one person must not
    # keep talking to Sunny after someone else signs in from another tab.
    page_user = request.headers.get("X-Sunny-User")
    if page_user and request.path.startswith("/api/voice/") and session.get("user_id") \
            and page_user != str(session.get("user_id")):
        return jsonify({"ok": False, "error": "Signed in as someone else", "account_changed": True}), 401

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        conn = get_conn(DB_PATH)
        user = login_user(conn, username, password)
        conn.close()
        if user:
            session.clear()
            session.permanent = True
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            return redirect(url_for("index"))
        return render_template("dashboard.html", view_mode="login", error="Invalid username or password")
    return render_template("dashboard.html", view_mode="login")

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        name = request.form.get("name", "").strip() or username.capitalize()
        designation = request.form.get("designation", "").strip() or "Product Manager"
        share_profile = 1 if request.form.get("share_profile") else 0
        linkedin_profile = request.form.get("linkedin_profile", "").strip()
        imap_email = request.form.get("imap_email", "").strip()
        imap_password = request.form.get("imap_password", "").strip()
        gemini_api_key = request.form.get("gemini_api_key", "").strip()
        
        # Check for uploaded resume file
        resume_text = ""
        resume_filename = ""
        file = request.files.get("resume_file")
        if file and file.filename:
            resume_filename = file.filename
            filename = file.filename.lower()
            if filename.endswith(".txt"):
                try:
                    resume_text = file.read().decode("utf-8", errors="ignore")
                except Exception:
                    pass
            elif filename.endswith(".pdf"):
                import pypdf
                try:
                    reader = pypdf.PdfReader(file)
                    resume_text = "\n".join([page.extract_text() or "" for page in reader.pages])
                except Exception as e:
                    print(f"Error parsing PDF: {e}")
                    resume_text = ""
            elif filename.endswith(".docx"):
                import docx
                try:
                    doc = docx.Document(file)
                    resume_text = "\n".join([p.text for p in doc.paragraphs])
                except Exception as e:
                    print(f"Error parsing DOCX: {e}")
                    resume_text = ""
        
        if not resume_text:
            resume_text = request.form.get("resume_text", "").strip()

        from resume_parser import parse_resume
        profile_json = ""
        if resume_text:
            try:
                profile_json = json.dumps(parse_resume(resume_text, gemini_api_key))
            except Exception as e:
                print(f"Error parsing resume: {e}")

        conn = get_conn(DB_PATH)
        try:
            uid = signup_user(conn, username, password)
            conn.execute(
                """UPDATE users SET resume_text = ?, imap_email = ?, imap_password = ?, gemini_api_key = ?, linkedin_profile = ?,
                   name = ?, designation = ?, share_profile = ?, resume_filename = ?, resume_profile_json = ?
                   WHERE id = ?""",
                (resume_text, imap_email, encrypt_secret(imap_password), encrypt_secret(gemini_api_key), linkedin_profile, name, designation, share_profile, resume_filename, profile_json, uid)
            )
            conn.commit()
            
            # If resume is provided on signup, scrape live jobs matching user designation and score them
            # (in the background, so signup doesn't hang on the scrapers)
            if resume_text:
                run_in_background(("sync", uid), run_user_sync, uid)

            session.clear()
            session.permanent = True
            session["user_id"] = uid
            session["username"] = username
            session["is_new_user"] = True
            conn.close()
            return redirect(url_for("index"))
        except Exception as e:
            conn.close()
            return render_template("dashboard.html", view_mode="signup", error=str(e))
    return render_template("dashboard.html", view_mode="signup")

def handle_social_login(provider, email, name):
    email = (email or "").strip().lower()
    if provider not in ("google", "linkedin") or not email or "@" not in email:
        return "Social authentication error: invalid profile", 400
    conn = get_conn(DB_PATH)
    try:
        # Match on the verified email address, not just its local part, so
        # jane@gmail.com and jane@company.com never share an account.
        row = conn.execute("SELECT id, username FROM users WHERE email = ? AND auth_provider = ?", (email, provider)).fetchone()
        if not row:
            # Accounts created before emails were stored used "<local part>_<provider>" as the
            # username; let the first verified login claim such a legacy account.
            legacy_username = email.split("@")[0] + "_" + provider
            row = conn.execute("SELECT id, username FROM users WHERE username = ? AND email IS NULL", (legacy_username,)).fetchone()
            if row:
                conn.execute("UPDATE users SET email = ?, auth_provider = ? WHERE id = ?", (email, provider, row[0]))
                conn.commit()
        if row:
            uid, username = row[0], row[1]
            is_new = False
        else:
            username = f"{email}_{provider}"
            password = secrets.token_urlsafe(24)
            uid = signup_user(conn, username, password, validate=False)
            # Profiles are only shown on the public landing page if the user opts in later
            conn.execute("UPDATE users SET name = ?, email = ?, auth_provider = ?, share_profile = 0 WHERE id = ?", (name, email, provider, uid))
            conn.commit()
            is_new = True

        session.clear()
        session.permanent = True
        session["user_id"] = uid
        session["username"] = username
        session["is_new_user"] = is_new
        conn.close()
        return redirect(url_for("index"))
    except Exception as e:
        conn.close()
        print(f"Social authentication error: {e}")
        return "Social authentication error. Please try again.", 500


@app.route("/auth/google")
def auth_google():
    from dotenv import load_dotenv
    base_dir = os.path.dirname(os.path.abspath(__file__))
    load_dotenv(dotenv_path=os.path.join(base_dir, ".env"), override=True)
    client_id = os.getenv("GOOGLE_CLIENT_ID")
    if not client_id:
        if ALLOW_DEV_LOGIN:
            return render_template("dashboard.html", view_mode="mock_auth", provider="google")
        return render_template("dashboard.html", view_mode="login", error="Google sign-in isn't set up on this server yet. Please use a username and password.")
    state = secrets.token_urlsafe(24)
    session["oauth_state"] = state
    redirect_uri = url_for("auth_google_callback", _external=True)
    google_auth_url = (
        f"https://accounts.google.com/o/oauth2/v2/auth?"
        f"client_id={client_id}&"
        f"redirect_uri={redirect_uri}&"
        f"response_type=code&"
        f"scope=openid%20email%20profile&"
        f"state={state}"
    )
    return redirect(google_auth_url)


@app.route("/auth/google/callback")
def auth_google_callback():
    from dotenv import load_dotenv
    base_dir = os.path.dirname(os.path.abspath(__file__))
    load_dotenv(dotenv_path=os.path.join(base_dir, ".env"), override=True)
    code = request.args.get("code")
    if not code:
        return "Authorization code missing", 400
    if not _check_oauth_state():
        return "Sign-in session expired or invalid. Please try again.", 400
    client_id = os.getenv("GOOGLE_CLIENT_ID")
    client_secret = os.getenv("GOOGLE_CLIENT_SECRET")
    redirect_uri = url_for("auth_google_callback", _external=True)
    
    import requests
    token_resp = requests.post(
        "https://oauth2.googleapis.com/token",
        data={
            "code": code,
            "client_id": client_id,
            "client_secret": client_secret,
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code"
        },
        timeout=15,
    )
    token_data = token_resp.json()
    access_token = token_data.get("access_token")
    if not access_token:
        print(f"OAuth token exchange failed: {token_data}")
        return "Failed to sign in. Please try again.", 400
        
    user_resp = requests.get(
        "https://www.googleapis.com/oauth2/v2/userinfo",
        headers={"Authorization": f"Bearer {access_token}"},
        timeout=15,
    )
    user_info = user_resp.json()
    email = user_info.get("email")
    if not email or user_info.get("verified_email") is False:
        return "Failed to retrieve a verified email from Google profile", 400
    name = user_info.get("name") or email.split("@")[0]
        
    return handle_social_login("google", email, name)


@app.route("/auth/linkedin")
def auth_linkedin():
    from dotenv import load_dotenv
    base_dir = os.path.dirname(os.path.abspath(__file__))
    load_dotenv(dotenv_path=os.path.join(base_dir, ".env"), override=True)
    client_id = os.getenv("LINKEDIN_CLIENT_ID")
    if not client_id:
        if ALLOW_DEV_LOGIN:
            return render_template("dashboard.html", view_mode="mock_auth", provider="linkedin")
        return render_template("dashboard.html", view_mode="login", error="LinkedIn sign-in isn't set up on this server yet. Please use a username and password.")
    state = secrets.token_urlsafe(24)
    session["oauth_state"] = state
    redirect_uri = url_for("auth_linkedin_callback", _external=True)
    linkedin_auth_url = (
        f"https://www.linkedin.com/oauth/v2/authorization?"
        f"client_id={client_id}&"
        f"redirect_uri={redirect_uri}&"
        f"response_type=code&"
        f"scope=openid%20profile%20email&"
        f"state={state}"
    )
    return redirect(linkedin_auth_url)


@app.route("/auth/linkedin/callback")
def auth_linkedin_callback():
    from dotenv import load_dotenv
    base_dir = os.path.dirname(os.path.abspath(__file__))
    load_dotenv(dotenv_path=os.path.join(base_dir, ".env"), override=True)
    code = request.args.get("code")
    if not code:
        return "Authorization code missing", 400
    if not _check_oauth_state():
        return "Sign-in session expired or invalid. Please try again.", 400
    client_id = os.getenv("LINKEDIN_CLIENT_ID")
    client_secret = os.getenv("LINKEDIN_CLIENT_SECRET")
    redirect_uri = url_for("auth_linkedin_callback", _external=True)
    
    import requests
    token_resp = requests.post(
        "https://www.linkedin.com/oauth/v2/accessToken",
        data={
            "code": code,
            "client_id": client_id,
            "client_secret": client_secret,
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code"
        },
        timeout=15,
    )
    token_data = token_resp.json()
    access_token = token_data.get("access_token")
    if not access_token:
        print(f"OAuth token exchange failed: {token_data}")
        return "Failed to sign in. Please try again.", 400
        
    user_resp = requests.get(
        "https://api.linkedin.com/v2/userinfo",
        headers={"Authorization": f"Bearer {access_token}"},
        timeout=15,
    )
    user_info = user_resp.json()
    email = user_info.get("email")
    if not email or user_info.get("email_verified") is False:
        return "Failed to retrieve a verified email from LinkedIn profile", 400
    name = user_info.get("name") or (user_info.get("given_name", "") + " " + user_info.get("family_name", "")).strip() or email.split("@")[0]
        
    return handle_social_login("linkedin", email, name)


@app.route("/auth/mock/callback", methods=["POST"])
def auth_mock_callback():
    if not ALLOW_DEV_LOGIN:
        abort(404)
    provider = request.form.get("provider")
    email = request.form.get("email")
    name = request.form.get("name")
    
    if not email or not name:
        return "Missing email or name", 400
        
    return handle_social_login(provider, email, name)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))

@app.route("/api/profile", methods=["POST"])
def update_profile():
    d = request.json
    uid = get_user_id()
    conn = get_conn(DB_PATH)
    
    # Fetch old values to check for changes
    old = conn.execute("SELECT resume_text, designation FROM users WHERE id = ?", (uid,)).fetchone()
    old_resume = old[0] if old else ""
    old_designation = old[1] if old else ""
    old_secrets = get_user_secrets(conn, uid)
    old_key = old_secrets["gemini_api_key"]
    
    new_resume = d.get("resume_text", "")
    new_key = d.get("gemini_api_key", "")
    if new_key == SECRET_PLACEHOLDER:
        new_key = old_key
    new_imap_password = d.get("imap_password", "")
    if new_imap_password == SECRET_PLACEHOLDER:
        new_imap_password = old_secrets["imap_password"]
    new_designation = d.get("designation", "")
    
    should_reparse = (new_resume != old_resume) or (new_key != old_key)
    designation_changed = (new_designation != old_designation)
    
    profile_json = None
    if should_reparse:
        from resume_parser import parse_resume
        try:
            profile_json = json.dumps(parse_resume(new_resume, new_key))
        except Exception as e:
            print(f"Error parsing resume: {e}")
    else:
        profile_row = conn.execute("SELECT resume_profile_json FROM users WHERE id = ?", (uid,)).fetchone()
        profile_json = profile_row[0] if profile_row else None
        
    conn.execute(
        """UPDATE users SET resume_text = ?, imap_email = ?, imap_password = ?, gemini_api_key = ?, linkedin_profile = ?,
           name = ?, designation = ?, share_profile = ?, resume_filename = ?, resume_profile_json = ?,
           weight_thumbs_up = ?, weight_applied = ?, weight_thumbs_down = ?, weight_rejected = ?,
           daily_recs_enabled = ?, daily_recs_min_score = ?, daily_recs_time = ?
           WHERE id = ?""",
        (new_resume, d.get("imap_email", ""), encrypt_secret(new_imap_password), encrypt_secret(new_key), d.get("linkedin_profile", ""),
         d.get("name", ""), d.get("designation", ""), d.get("share_profile", 0), d.get("resume_filename", ""), profile_json,
         float(d.get("weight_thumbs_up", 1.0)), float(d.get("weight_applied", 1.0)), float(d.get("weight_thumbs_down", -1.0)), float(d.get("weight_rejected", -1.5)),
         int(d.get("daily_recs_enabled", 1)), float(d.get("daily_recs_min_score", 7.5)), d.get("daily_recs_time", "07:30"), uid)
    )
    conn.commit()
    
    if designation_changed:
        conn.execute("DELETE FROM jobs WHERE user_id = ? AND status = 'new'", (uid,))
        conn.commit()
        
    conn.close()
    if should_reparse:
        rescore_inbox_in_background(uid)
    return jsonify({"ok": True})


@app.route("/api/jobs/feedback", methods=["POST"])
def update_job_feedback():
    uid = get_user_id()
    d = request.json or {}
    job_id = d.get("job_id")
    feedback = d.get("feedback") # 1, -1, or 0
    
    if not job_id or feedback is None:
        return jsonify({"ok": False, "error": "Missing parameters"}), 400
        
    conn = get_conn(DB_PATH)
    conn.execute(
        "UPDATE jobs SET feedback = ? WHERE job_id = ? AND user_id = ?",
        (feedback, job_id, uid)
    )
    conn.commit()
    conn.close()
    
    # Re-score the inbox with the new preference signal, in the background
    rescore_inbox_in_background(uid)
    return jsonify({"ok": True})

@app.route("/api/profile/upload", methods=["POST"])
def upload_profile_resume():
    uid = get_user_id()
    file = request.files.get("resume_file")
    if not file or not file.filename:
        return jsonify({"ok": False, "error": "No file uploaded"}), 400
        
    filename = file.filename.lower()
    resume_text = ""
    if filename.endswith(".txt"):
        resume_text = file.read().decode("utf-8", errors="ignore")
    elif filename.endswith(".pdf"):
        import pypdf
        try:
            reader = pypdf.PdfReader(file)
            resume_text = "\n".join([page.extract_text() or "" for page in reader.pages])
        except Exception as e:
            return jsonify({"ok": False, "error": f"Error parsing PDF: {str(e)}"}), 400
    elif filename.endswith(".docx"):
        import docx
        try:
            doc = docx.Document(file)
            resume_text = "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            return jsonify({"ok": False, "error": f"Error parsing DOCX: {str(e)}"}), 400
    else:
        return jsonify({"ok": False, "error": "Unsupported file format. Please upload PDF, TXT, or DOCX."}), 400
        
    conn = get_conn(DB_PATH)
    api_key = get_user_secrets(conn, uid)["gemini_api_key"] or None
    
    from resume_parser import parse_resume
    profile_json = ""
    try:
        profile_json = json.dumps(parse_resume(resume_text, api_key))
    except Exception as e:
        print(f"Error parsing resume: {e}")
        
    conn.execute("UPDATE users SET resume_text = ?, resume_filename = ?, resume_profile_json = ? WHERE id = ?", (resume_text, file.filename, profile_json, uid))
    conn.commit()
    conn.close()
    
    # Re-score in the background to prevent gateway timeouts on large pipelines
    rescore_inbox_in_background(uid)
    return jsonify({"ok": True, "resume_text": resume_text, "resume_filename": file.filename})


@app.route("/api/onboard_skip", methods=["POST"])
def onboard_skip():
    session["skip_onboarding"] = True
    return jsonify({"ok": True})


@app.route("/api/onboard_resume", methods=["POST"])
def onboard_resume():
    uid = get_user_id()
    file = request.files.get("resume_file")
    
    resume_text = ""
    resume_filename = "Pasted_Resume.txt"
    if file and file.filename:
        resume_filename = file.filename
        filename = file.filename.lower()
        if filename.endswith(".txt"):
            resume_text = file.read().decode("utf-8", errors="ignore")
        elif filename.endswith(".pdf"):
            import pypdf
            try:
                reader = pypdf.PdfReader(file)
                resume_text = "\n".join([page.extract_text() or "" for page in reader.pages])
            except Exception as e:
                return jsonify({"ok": False, "error": f"Error parsing PDF: {str(e)}"}), 400
        elif filename.endswith(".docx"):
            import docx
            try:
                doc = docx.Document(file)
                resume_text = "\n".join([p.text for p in doc.paragraphs])
            except Exception as e:
                return jsonify({"ok": False, "error": f"Error parsing DOCX: {str(e)}"}), 400
        else:
            return jsonify({"ok": False, "error": "Unsupported file format. Please upload PDF, TXT, or DOCX."}), 400
    else:
        # Fallback to form field or JSON body
        resume_text = request.form.get("resume_text", "").strip()
        if not resume_text:
            try:
                d = request.json or {}
                resume_text = d.get("resume_text", "").strip()
            except Exception:
                pass

    if not resume_text:
        return jsonify({"ok": False, "error": "Resume text is empty"}), 400

    conn = get_conn(DB_PATH)
    api_key = get_user_secrets(conn, uid)["gemini_api_key"] or None
    
    from resume_parser import parse_resume
    profile_json = ""
    try:
        profile_json = json.dumps(parse_resume(resume_text, api_key))
    except Exception as e:
        print(f"Error parsing resume: {e}")

    conn.execute("UPDATE users SET resume_text = ?, resume_filename = ?, resume_profile_json = ? WHERE id = ?", (resume_text, resume_filename, profile_json, uid))
    conn.commit()

    conn.close()

    # Scraping + scoring can take minutes (and longer with AI scoring), which used to
    # outlast the server's request timeout and return an HTML error page. Run it in the
    # background and wait only briefly, so the user sees early results if there are any.
    import time
    run_in_background(("sync", uid), run_user_sync, uid)
    deadline = time.time() + ONBOARD_WAIT_SECONDS
    while time.time() < deadline:
        with _running_lock:
            if ("sync", uid) not in _running_jobs:
                break
        time.sleep(0.5)
    with _running_lock:
        still_syncing = ("sync", uid) in _running_jobs

    conn = get_conn(DB_PATH)
    suggestions = conn.execute(
        """SELECT title, company, location, ai_score, ai_summary 
           FROM jobs WHERE user_id = ? AND ai_score IS NOT NULL 
           ORDER BY ai_score DESC LIMIT 3""", (uid,)
    ).fetchall()
    
    conn.close()
    return jsonify({
        "ok": True,
        "suggestions": [dict(s) for s in suggestions],
        "still_syncing": still_syncing
    })


@app.route("/api/resume/parse", methods=["POST"])
def api_resume_parse():
    file = request.files.get("resume_file")
    resume_text = ""
    if file and file.filename:
        filename = file.filename.lower()
        if filename.endswith(".txt"):
            resume_text = file.read().decode("utf-8", errors="ignore")
        elif filename.endswith(".pdf"):
            import pypdf
            try:
                reader = pypdf.PdfReader(file)
                resume_text = "\n".join([page.extract_text() or "" for page in reader.pages])
            except Exception as e:
                return jsonify({"ok": False, "error": f"Error parsing PDF: {str(e)}"}), 400
        elif filename.endswith(".docx"):
            import docx
            try:
                doc = docx.Document(file)
                resume_text = "\n".join([p.text for p in doc.paragraphs])
            except Exception as e:
                return jsonify({"ok": False, "error": f"Error parsing DOCX: {str(e)}"}), 400
        else:
            return jsonify({"ok": False, "error": "Unsupported file format"}), 400
    else:
        try:
            d = request.json or {}
            resume_text = d.get("resume_text", "").strip()
        except Exception:
            resume_text = request.form.get("resume_text", "").strip()
            
    if not resume_text:
        return jsonify({"ok": False, "error": "No resume content provided"}), 400
        
    from resume_parser import parse_resume
    try:
        uid = get_user_id()
        conn = get_conn(DB_PATH)
        api_key = get_user_secrets(conn, uid)["gemini_api_key"] or None
        conn.close()
        
        profile = parse_resume(resume_text, api_key)
        return jsonify({"ok": True, "profile": profile})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ── Helpers ────────────────────────────────────────────────────────────────

def _stats(conn):
    uid = get_user_id()
    total = conn.execute("SELECT COUNT(*) FROM jobs WHERE user_id = ?", (uid,)).fetchone()[0]
    by_status = dict(conn.execute("SELECT status, COUNT(*) FROM jobs WHERE user_id = ? GROUP BY status", (uid,)).fetchall())
    avg = conn.execute("SELECT AVG(ai_score) FROM jobs WHERE user_id = ? AND ai_score IS NOT NULL", (uid,)).fetchone()[0]
    email_count = conn.execute("SELECT COUNT(*) FROM received_emails WHERE user_id = ?", (uid,)).fetchone()[0]
    
    last_scraped_row = conn.execute("SELECT last_scraped_at FROM users WHERE id = ?", (uid,)).fetchone()
    last_updated = last_scraped_row[0] if last_scraped_row and last_scraped_row[0] else None
    
    if not last_updated:
        last_updated_row = conn.execute("SELECT MAX(scraped_at) FROM jobs WHERE user_id = ?", (uid,)).fetchone()
        last_updated = last_updated_row[0] if last_updated_row and last_updated_row[0] else None
        
    if last_updated:
        try:
            dt = datetime.fromisoformat(last_updated.split('.')[0])
            last_updated_str = dt.strftime("%b %d, %I:%M %p")
        except Exception:
            last_updated_str = last_updated
    else:
        last_updated_str = "Never"
        
    return {"total": total, "by_status": by_status,
            "avg_score": round(avg, 1) if avg else 0,
            "last_updated": last_updated_str,
            "email_count": email_count}


def _batch_full_jobs(conn, raw_jobs, uid):
    if not raw_jobs:
        return []

    # 1. Cover letters
    cl_rows = conn.execute("SELECT * FROM cover_letters WHERE user_id = ?", (uid,)).fetchall()
    cls = {row["job_id"]: dict(row) for row in cl_rows}

    # 2. Contacts
    contact_rows = conn.execute("SELECT * FROM contacts WHERE user_id = ?", (uid,)).fetchall()
    contacts = {}
    for row in contact_rows:
        jid = row["job_id"]
        if jid not in contacts:
            contacts[jid] = []
        contacts[jid].append(dict(row))

    # 3. Notes
    notes_rows = conn.execute("SELECT * FROM application_notes WHERE user_id = ?", (uid,)).fetchall()
    notes = {row["job_id"]: dict(row) for row in notes_rows}

    # 4. Timeline
    timeline_rows = conn.execute("SELECT job_id, event, created_at FROM application_timeline WHERE user_id = ? ORDER BY created_at", (uid,)).fetchall()
    timelines = {}
    for row in timeline_rows:
        jid = row["job_id"]
        if jid not in timelines:
            timelines[jid] = []
        timelines[jid].append({"event": row["event"], "created_at": row["created_at"]})

    # 5. Emails
    email_rows = conn.execute("SELECT * FROM received_emails WHERE user_id = ? ORDER BY received_at DESC", (uid,)).fetchall()
    emails = {}
    for row in email_rows:
        jid = row["job_id"]
        if jid not in emails:
            emails[jid] = []
        emails[jid].append(dict(row))

    # Construct the full job objects
    full_jobs = []
    for row in raw_jobs:
        j = dict(row)
        jid = j["job_id"]
        
        j["key_reqs_list"] = json.loads(j.get("key_reqs") or "[]")
        j["matched_skills_list"] = json.loads(j.get("matched_skills") or "[]")
        j["missing_skills_list"] = json.loads(j.get("missing_skills") or "[]")
        
        j["cover_letter"] = cls.get(jid)
        j["contacts"] = contacts.get(jid, [])
        j["notes"] = notes.get(jid, {"note": "", "linkedin_note": ""})
        j["timeline"] = timelines.get(jid, [])
        j["emails"] = emails.get(jid, [])
        
        full_jobs.append(j)
        
    return full_jobs



# ── Pages ──────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    conn = get_conn(DB_PATH)
    uid = session.get("user_id")
    print(f"DEBUG: index route called. session user_id = {uid}")
    if not uid:
        # Fetch consenting seekers for landing page wall (excluding admin)
        users = conn.execute(
            "SELECT name, designation, linkedin_profile FROM users WHERE share_profile = 1 AND username != 'admin' ORDER BY id DESC"
        ).fetchall()
        conn.close()
        return render_template("landing.html", seekers=[dict(u) for u in users])
        
    # Get user settings to pass to frontend profile form (stored secrets are never sent back)
    settings = get_user_settings(conn, uid, mask_secrets=True)
    needs_onboarding = not settings.get("resume_text") and not session.get("skip_onboarding")
    if needs_onboarding:
        stats = _stats(conn)
        conn.close()
        return render_template("dashboard.html", settings=settings, view_mode="onboarding", stats=stats)
        
    # Fetch all active jobs (ignore archived)
    raw_jobs = conn.execute("SELECT * FROM jobs WHERE user_id = ? AND status != 'archived' ORDER BY COALESCE(ai_score,0) DESC", (uid,)).fetchall()
    
    board = {
        "whatsapp": [],
        "new": [],
        "shortlisted": [],
        "interviewing": [],
        "applied": [],
        "offer": [],
        "rejected": []
    }
    
    active_jobs = _batch_full_jobs(conn, raw_jobs, uid)
    for j in active_jobs:
        status = j["status"]
        if status == "whatsapp":
            board["whatsapp"].append(j)
        elif status in ("new", "scored", "ready"):
            board["new"].append(j)
        elif status == "shortlisted":
            board["shortlisted"].append(j)
        elif status == "interviewing":
            board["interviewing"].append(j)
        elif status == "applied":
            board["applied"].append(j)
        elif status == "offer":
            board["offer"].append(j)
        elif status == "rejected":
            board["rejected"].append(j)
            
    stats = _stats(conn)
    companies = [r[0] for r in conn.execute(
        "SELECT DISTINCT company FROM jobs WHERE user_id = ? ORDER BY company", (uid,)).fetchall()]
    conn.close()
    
    is_new_user = session.get("is_new_user", False)
    if is_new_user:
        session.pop("is_new_user", None)
        
    cols = ["whatsapp", "new", "shortlisted", "interviewing", "applied", "offer", "rejected"]
    return render_template("dashboard.html", board=board, cols=cols, stats=stats,
                           companies=companies, settings=settings, view_mode="board", is_new_user=is_new_user)


@app.route("/emails")
def view_emails():
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    raw_emails = conn.execute(
        """SELECT r.*, j.company, j.title 
           FROM received_emails r 
           LEFT JOIN jobs j ON r.job_id = j.job_id 
           WHERE r.user_id = ?
           ORDER BY r.received_at DESC""", (uid,)
    ).fetchall()
    emails = [dict(e) for e in raw_emails]
    stats = _stats(conn)
    conn.close()
    
    return render_template("dashboard.html", view_mode="emails", emails=emails, stats=stats)


@app.route("/pipeline")
def pipeline():
    return redirect("/")


@app.route("/users", methods=["GET", "POST"])
def admin_users():
    if not is_admin():
        return redirect("/")
    
    conn = get_conn(DB_PATH)
    
    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if not username or not password:
            error = "Both fields are required."
        else:
            existing = conn.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone()
            if existing:
                error = "Username already exists."
            else:
                from datetime import datetime
                from auth import hash_password
                from auth import validate_new_credentials
                try:
                    validate_new_credentials(username.lower(), password)
                    conn.execute(
                        "INSERT INTO users (username, password_hash, created_at) VALUES (?, ?, ?)",
                        (username.lower(), hash_password(password), datetime.now().isoformat())
                    )
                    conn.commit()
                except ValueError as e:
                    error = str(e)
                
    # Fetch all users and calculate their job counts
    raw_users = conn.execute("SELECT id, username, name, created_at FROM users ORDER BY username").fetchall()
    users = []
    for u in raw_users:
        u_dict = dict(u)
        job_count = conn.execute("SELECT COUNT(*) FROM jobs WHERE user_id = ?", (u_dict["id"],)).fetchone()[0]
        u_dict["job_count"] = job_count
        users.append(u_dict)
        
    stats = _stats(conn)
    conn.close()
    return render_template("dashboard.html", view_mode="users", users=users, stats=stats, error=error)


@app.route("/api/users/<int:user_id>/delete", methods=["POST"])
def delete_user(user_id):
    if not is_admin():
        return jsonify({"error": "unauthorized"}), 403
    if user_id == 1:
        return jsonify({"error": "cannot delete admin"}), 400
        
    conn = get_conn(DB_PATH)
    conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.execute("DELETE FROM jobs WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM contacts WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM cover_letters WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM application_notes WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM application_timeline WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM received_emails WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM tailored_resumes WHERE user_id = ?", (user_id,))
    conn.execute("DELETE FROM interview_prep WHERE user_id = ?", (user_id,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# ── API ────────────────────────────────────────────────────────────────────

@app.route("/api/job/<job_id>/status", methods=["POST"])
def set_status(job_id):
    status = request.json.get("status")
    allowed = {"whatsapp", "new", "scored", "shortlisted", "ready", "applied", "interviewing", "offer", "rejected", "archived"}
    if status not in allowed:
        return jsonify({"error": "invalid"}), 400
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    conn.execute("UPDATE jobs SET status=? WHERE job_id=? AND user_id=?", (status, job_id, uid))
    add_timeline(conn, job_id, f"Status → {status}")
    conn.commit()
    conn.close()
    
    # Re-score the inbox with the new preference signal, in the background
    rescore_inbox_in_background(uid)
    return jsonify({"ok": True, "status": status})


@app.route("/api/job/<job_id>/cover-letter", methods=["PUT"])
def save_cover_letter(job_id):
    d = request.json
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    conn.execute("UPDATE cover_letters SET subject=?, body=? WHERE job_id=? AND user_id=?",
                 (d["subject"], d["body"], job_id, uid))
    conn.commit(); conn.close()
    return jsonify({"ok": True})


@app.route("/api/job/<job_id>/note", methods=["POST"])
def save_note(job_id):
    d = request.json
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    conn.execute(
        "INSERT OR REPLACE INTO application_notes (job_id, note, linkedin_note, user_id) VALUES (?,?,?,?)",
        (job_id, d.get("note",""), d.get("linkedin_note",""), uid)
    )
    conn.commit(); conn.close()
    return jsonify({"ok": True})


@app.route("/api/job/<job_id>/regenerate", methods=["POST"])
def regenerate(job_id):
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    job = conn.execute("SELECT * FROM jobs WHERE job_id=? AND user_id=?", (job_id, uid)).fetchone()
    if not job:
        conn.close(); return jsonify({"error": "not found"}), 404
    j = dict(job)
    contact = conn.execute(
        "SELECT name, title FROM contacts WHERE job_id=? AND user_id=? LIMIT 1", (job_id, uid)
    ).fetchone()
    cn = contact[0] if contact else "Hiring Team"
    ct = contact[1] if contact else "Recruiter"

    settings = get_user_settings(conn, uid)
    score_data = score_job(j["title"], j["company"], j.get("description",""), resume_text=settings.get("resume_text"), api_key=settings.get("gemini_api_key"), user_id=uid)
    letter     = generate_cover_letter(j["title"], j["company"], j.get("description",""), cn, ct, resume_text=settings.get("resume_text"), api_key=settings.get("gemini_api_key"), candidate_name=settings.get("name") or "")
    li_note    = generate_linkedin_note(cn, ct, j["company"], j["title"], api_key=settings.get("gemini_api_key"))

    conn.execute("UPDATE jobs SET ai_score=?, ai_summary=?, key_reqs=? WHERE job_id=? AND user_id=?",
                 (score_data["score"], score_data["fit_summary"],
                  json.dumps(score_data.get("key_requirements",[])), job_id, uid))
    conn.execute(
        "INSERT OR REPLACE INTO cover_letters (job_id, subject, body, linkedin_note, created_at, user_id) VALUES (?,?,?,?,?,?)",
        (job_id, letter["subject"], letter["body"], li_note, datetime.now().isoformat(), uid)
    )
    add_timeline(conn, job_id, "Regenerated cover letter")
    conn.commit(); conn.close()
    return jsonify({"ok": True, "score": score_data["score"],
                    "fit_summary": score_data["fit_summary"],
                    "cover_letter": letter, "linkedin_note": li_note,
                    "key_requirements": score_data.get("key_requirements",[])})


@app.route("/api/job/<job_id>/send-email", methods=["POST"])
def send_email(job_id):
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    j_row = conn.execute("SELECT * FROM jobs WHERE job_id=? AND user_id=?", (job_id, uid)).fetchone()
    if not j_row:
        conn.close(); return jsonify({"error": "not found"}), 404
    j  = dict(j_row)
    cl = conn.execute("SELECT * FROM cover_letters WHERE job_id=? AND user_id=?", (job_id, uid)).fetchone()
    c  = conn.execute("SELECT * FROM contacts WHERE job_id=? AND user_id=? LIMIT 1", (job_id, uid)).fetchone()
    conn.close()
    if not cl:
        return jsonify({"error": "no cover letter"}), 400
    item = {**j, "score": j.get("ai_score",0),
            "fit_summary": j.get("ai_summary",""),
            "key_requirements": json.loads(j.get("key_reqs") or "[]"),
            "cover_letter_subject": cl["subject"],
            "cover_letter_body": cl["body"],
            "contact_name": c["name"] if c else "Hiring Team",
            "contact_title": c["title"] if c else "",
            "linkedin_note": cl.get("linkedin_note","") if cl else ""}
    ok = send_email_digest([item], recipient=recipient_for_user(uid))
    if ok:
        conn2 = get_conn(DB_PATH)
        conn2.execute("UPDATE jobs SET status='applied' WHERE job_id=? AND user_id=?", (job_id, uid))
        add_timeline(conn2, job_id, "Email sent → applied")
        conn2.commit(); conn2.close()
    return jsonify({"ok": ok, "message": "Email sent!" if ok else "Couldn't send. Add your email address in settings, or email isn't configured on this server."})


@app.route("/api/bulk", methods=["POST"])
def bulk_action():
    d      = request.json
    action = d.get("action")
    ids    = d.get("job_ids", [])
    conn   = get_conn(DB_PATH)
    uid    = get_user_id()

    if action in ("shortlist", "archive", "applied"):
        status_map = {"shortlist": "shortlisted", "archive": "archived", "applied": "applied"}
        for jid in ids:
            conn.execute("UPDATE jobs SET status=? WHERE job_id=? AND user_id=?", (status_map[action], jid, uid))
            add_timeline(conn, jid, f"Bulk → {status_map[action]}")

    elif action == "send-digest":
        items = []
        for jid in ids:
            j  = conn.execute("SELECT * FROM jobs WHERE job_id=? AND user_id=?", (jid, uid)).fetchone()
            cl = conn.execute("SELECT * FROM cover_letters WHERE job_id=? AND user_id=?", (jid, uid)).fetchone()
            c  = conn.execute("SELECT * FROM contacts WHERE job_id=? AND user_id=? LIMIT 1", (jid, uid)).fetchone()
            if j and cl:
                jd = dict(j)
                items.append({
                    **jd, "score": jd.get("ai_score",0),
                    "fit_summary": jd.get("ai_summary",""),
                    "key_requirements": json.loads(jd.get("key_reqs") or "[]"),
                    "cover_letter_subject": cl["subject"],
                    "cover_letter_body": cl["body"],
                    "contact_name": c["name"] if c else "Hiring Team",
                    "contact_title": c["title"] if c else "",
                    "linkedin_note": cl.get("linkedin_note","") if cl else "",
                })
        send_email_digest(items, recipient=recipient_for_user(uid))

    conn.commit(); conn.close()
    return jsonify({"ok": True})


@app.route("/api/pipeline/move", methods=["POST"])
def pipeline_move():
    d = request.json
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    conn.execute("UPDATE jobs SET status=? WHERE job_id=? AND user_id=?", (d["status"], d["job_id"], uid))
    add_timeline(conn, d["job_id"], f"Pipeline → {d['status']}")
    conn.commit(); conn.close()
    return jsonify({"ok": True})


def run_daily_recommendations():
    conn = get_conn(DB_PATH)
    users = conn.execute("SELECT id, daily_recs_min_score FROM users WHERE daily_recs_enabled = 1").fetchall()
    conn.close()

    for user in users:
        uid = user[0]
        min_score = user[1] if user[1] is not None else 7.5
        try:
            # 1. Clear old daily picks and remember which jobs already existed
            conn = get_conn(DB_PATH)
            conn.execute("UPDATE jobs SET is_daily_pick = 0 WHERE user_id = ?", (uid,))
            conn.commit()
            before = {r[0] for r in conn.execute("SELECT job_id FROM jobs WHERE user_id = ?", (uid,)).fetchall()}
            conn.close()

            # 2. Scrape, enrich, sync email, score
            run_user_sync(uid, min_score=0)

            # 3. Today's picks: jobs added in this run that clear the user's bar
            conn = get_conn(DB_PATH)
            rows = conn.execute(
                "SELECT job_id, title, company, location, url, ai_score, ai_summary FROM jobs WHERE user_id = ? AND ai_score >= ?",
                (uid, min_score)
            ).fetchall()
            picks = [dict(r) for r in rows if r[0] not in before]
            now = datetime.now().isoformat()
            for pick in picks:
                conn.execute("UPDATE jobs SET is_daily_pick = 1, picked_at = ? WHERE job_id = ? AND user_id = ?", (now, pick["job_id"], uid))
            conn.commit()
            conn.close()

            # 4. Email digest to this user only
            to = recipient_for_user(uid)
            if picks and to:
                send_email_digest([{**pick, "score": pick.get("ai_score") or 0, "fit_summary": pick.get("ai_summary") or ""} for pick in picks], recipient=to)
        except Exception as e:
            print(f"Daily recommendations failed for user {uid}: {e}")


@app.route("/api/cron/daily-recommendations", methods=["POST"])
def cron_daily_recommendations():
    expected_token = os.environ.get("CRON_SECRET", "")
    if not expected_token:
        return jsonify({"ok": False, "error": "CRON_SECRET is not configured"}), 503
    auth_header = request.headers.get("X-Cron-Token", "")
    if not hmac.compare_digest(auth_header, expected_token):
        return jsonify({"ok": False, "error": "Unauthorized"}), 401

    started = run_in_background(("daily-recs",), run_daily_recommendations)
    return jsonify({"ok": True, "started": started})


@app.route("/api/recommendations/today", methods=["GET"])
def get_daily_recommendations():
    uid = get_user_id()
    conn = get_conn(DB_PATH)
    picks = conn.execute(
        """SELECT job_id, title, company, location, url, ai_score, ai_summary, picked_at 
           FROM jobs 
           WHERE user_id = ? AND is_daily_pick = 1 
           ORDER BY ai_score DESC""", (uid,)
    ).fetchall()
    conn.close()
    return jsonify({
        "ok": True,
        "picks": [dict(p) for p in picks]
    })


@app.route("/api/job/<job_id>/interview-round", methods=["POST"])
def update_interview_round(job_id):
    uid = get_user_id()
    d = request.json or {}
    round_name = d.get("round", "").strip()
    
    conn = get_conn(DB_PATH)
    row = conn.execute("SELECT 1 FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
    if not row:
        conn.close()
        return jsonify({"ok": False, "error": "Job not found"}), 404
        
    from datetime import datetime
    
    conn.execute(
        "UPDATE jobs SET interview_round = ?, interview_round_updated_at = ? WHERE job_id = ? AND user_id = ?",
        (round_name, datetime.now().isoformat(), job_id, uid)
    )
    conn.commit()
    
    add_timeline(conn, job_id, f"Interview round -> {round_name}")
    
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/job/<job_id>/interview-prep", methods=["GET"])
def get_job_interview_prep(job_id):
    uid = get_user_id()
    conn = get_conn(DB_PATH)
    
    prep = conn.execute(
        "SELECT quick_questions, deep_questions FROM interview_prep WHERE job_id = ? AND user_id = ?",
        (job_id, uid)
    ).fetchone()
    
    if prep:
        conn.close()
        return jsonify({
            "ok": True,
            "quick_questions": json.loads(prep[0] or "[]"),
            "deep_questions": json.loads(prep[1] or "[]")
        })
        
    job = conn.execute(
        "SELECT title, company, description FROM jobs WHERE job_id = ? AND user_id = ?",
        (job_id, uid)
    ).fetchone()
    
    if not job:
        conn.close()
        return jsonify({"ok": False, "error": "Job not found"}), 404
        
    row = conn.execute("SELECT resume_text FROM users WHERE id = ?", (uid,)).fetchone()
    resume_text = row[0] if row else None
    api_key = get_user_secrets(conn, uid)["gemini_api_key"] or None
    
    from ai_engine import generate_interview_prep
    from datetime import datetime
    
    try:
        prep_data = generate_interview_prep(job[0], job[1], job[2] or "", resume_text, api_key)
        quick_json = json.dumps(prep_data.get("quick_questions", []))
        deep_json = json.dumps(prep_data.get("deep_questions", []))
        
        conn.execute(
            """INSERT OR REPLACE INTO interview_prep (job_id, quick_questions, deep_questions, created_at, user_id)
               VALUES (?, ?, ?, ?, ?)""",
            (job_id, quick_json, deep_json, datetime.now().isoformat(), uid)
        )
        conn.commit()
        conn.close()
        
        return jsonify({
            "ok": True,
            "quick_questions": prep_data.get("quick_questions", []),
            "deep_questions": prep_data.get("deep_questions", [])
        })
    except Exception as e:
        conn.close()
        return jsonify({"ok": False, "error": f"Failed to generate prep questions: {e}"}), 500


@app.route("/api/job/<job_id>/interview-prep/feedback", methods=["POST"])
def interview_prep_feedback(job_id):
    """Grades one practice answer (spoken or typed) for a job the user owns."""
    uid = get_user_id()
    d = request.get_json(silent=True) or {}
    g.voice_lang = ve.reply_lang(d.get("lang"), str(d.get("answer") or ""))
    limited = _voice_rate_limited(uid, "quiz_feedback", 20)
    if limited:
        return limited
    question = str(d.get("question") or "").strip()
    if not question:
        return jsonify({"ok": False, "error": "No question provided"}), 400
    conn = get_conn(DB_PATH)
    try:
        job = conn.execute("SELECT title, company FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
        if not job:
            return jsonify({"ok": False, "error": "Job not found"}), 404
        api_key = _voice_api_key(conn, uid)
    finally:
        conn.close()
    result = quiz_coach.grade_answer(question, d.get("hints") or "", d.get("answer") or "",
                                     title=job[0] or "", company=job[1] or "", api_key=api_key, lang=g.voice_lang)
    return jsonify({"ok": True, "lang": g.voice_lang, **result})


@app.route("/api/job/<job_id>/interview-prep/regenerate", methods=["POST"])
def regenerate_job_interview_prep(job_id):
    uid = get_user_id()
    conn = get_conn(DB_PATH)
    
    job = conn.execute(
        "SELECT title, company, description FROM jobs WHERE job_id = ? AND user_id = ?",
        (job_id, uid)
    ).fetchone()
    
    if not job:
        conn.close()
        return jsonify({"ok": False, "error": "Job not found"}), 404
        
    row = conn.execute("SELECT resume_text FROM users WHERE id = ?", (uid,)).fetchone()
    resume_text = row[0] if row else None
    api_key = get_user_secrets(conn, uid)["gemini_api_key"] or None
    
    from ai_engine import generate_interview_prep
    from datetime import datetime
    
    try:
        prep_data = generate_interview_prep(job[0], job[1], job[2] or "", resume_text, api_key)
        quick_json = json.dumps(prep_data.get("quick_questions", []))
        deep_json = json.dumps(prep_data.get("deep_questions", []))
        
        conn.execute(
            """INSERT OR REPLACE INTO interview_prep (job_id, quick_questions, deep_questions, created_at, user_id)
               VALUES (?, ?, ?, ?, ?)""",
            (job_id, quick_json, deep_json, datetime.now().isoformat(), uid)
        )
        conn.commit()
        conn.close()
        
        return jsonify({
            "ok": True,
            "quick_questions": prep_data.get("quick_questions", []),
            "deep_questions": prep_data.get("deep_questions", [])
        })
    except Exception as e:
        conn.close()
        return jsonify({"ok": False, "error": f"Failed to regenerate prep questions: {e}"}), 500


@app.route("/api/stats")
def api_stats():
    conn = get_conn(DB_PATH)
    s = _stats(conn)
    conn.close()
    return jsonify(s)


def run_user_sync(uid, min_score=None):
    """Scrape → contacts → email sync → score for one user. Every step is scoped to uid."""
    from scraper import run_all_scrapers
    from linkedin_finder import enrich_jobs_with_contacts
    from email_scraper import sync_job_statuses_from_email
    from ai_engine import process_new_jobs

    print(f"Background sync started for user {uid}")
    run_all_scrapers(DB_PATH, user_id=uid)
    enrich_jobs_with_contacts(DB_PATH, user_id=uid)
    try:
        sync_job_statuses_from_email(DB_PATH, user_id=uid)
    except Exception as e:
        print(f"Email sync failed for user {uid}: {e}")
    if min_score is None:
        min_score = float(os.getenv("MIN_SCORE", "0.0"))
    digest = process_new_jobs(DB_PATH, min_score=min_score, user_id=uid)

    conn = get_conn(DB_PATH)
    conn.execute("UPDATE users SET last_scraped_at = ? WHERE id = ?", (datetime.now().isoformat(), uid))
    conn.commit()
    conn.close()
    print(f"Background sync completed for user {uid}")
    return digest


@app.route("/api/refresh", methods=["POST"])
def refresh_listings():
    uid = get_user_id()
    started = run_in_background(("sync", uid), run_user_sync, uid)
    return jsonify({"ok": True, "message": "Sync started in background" if started else "A sync is already running"})


# ── Web job search (the board's search bar) ───────────────────────────────

_LINKEDIN_JOB_URL = re.compile(r"^https://([a-z]{2,3}\.)?linkedin\.com/jobs/view/[\w%-]+/?$")


@app.route("/api/jobs/search", methods=["GET"])
def search_web_jobs_api():
    from scraper import search_web_jobs
    uid = get_user_id()
    query = (request.args.get("q") or "").strip()[:80]
    location = (request.args.get("location") or "").strip()[:60]
    if len(query) < 2:
        return jsonify({"ok": False, "error": "Type at least two letters to search."}), 400
    if not ve.rate_limiter.allow(uid, "web_search", 10):
        return jsonify({"ok": False, "error": "That's a lot of searches. Give it a minute and try again."}), 429

    found = search_web_jobs(query, location)
    conn = get_conn(DB_PATH)
    rows = conn.execute("SELECT url, LOWER(title), LOWER(company) FROM jobs WHERE user_id = ?", (uid,)).fetchall()
    conn.close()
    on_board_urls = {r[0] for r in rows if r[0]}
    on_board_pairs = {(r[1], r[2]) for r in rows}
    for r in found["results"]:
        r["on_board"] = r["url"] in on_board_urls or (r["title"].lower(), r["company"].lower()) in on_board_pairs
    return jsonify({"ok": True, "query": query, "location": location, **found})


def _score_added_job(uid, job_id):
    conn = get_conn(DB_PATH)
    try:
        job = conn.execute("SELECT title, company, description FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
        if not job:
            return
        settings = get_user_settings(conn, uid)
        data = score_job(job["title"], job["company"], job["description"] or "", resume_text=settings.get("resume_text"),
                         api_key=settings.get("gemini_api_key"), user_id=uid)
        conn.execute("UPDATE jobs SET ai_score = ?, ai_summary = ?, key_reqs = ?, status = 'scored' WHERE job_id = ? AND user_id = ?",
                     (data.get("score"), data.get("fit_summary") or data.get("summary", ""),
                      json.dumps(data.get("key_requirements", [])), job_id, uid))
        conn.commit()
    finally:
        conn.close()


@app.route("/api/jobs/search/add", methods=["POST"])
def add_web_search_result():
    from scraper import find_board_posting, linkedin_job_id, _fetch_description, _insert_job
    uid = get_user_id()
    data = request.get_json(silent=True) or {}
    source = data.get("source")

    if source == "direct":
        try:
            job = find_board_posting(str(data.get("ref", "")))
        except Exception as e:
            print(f"[web search add] {e}")
            job = None
        if not job:
            return jsonify({"ok": False, "error": "That opening isn't listed anymore."}), 404
    elif source == "linkedin":
        url = str(data.get("ref", "")).strip()
        if not _LINKEDIN_JOB_URL.match(url):
            return jsonify({"ok": False, "error": "That doesn't look like a LinkedIn job link."}), 400
        title = str(data.get("title", "")).strip()[:200]
        company = str(data.get("company", "")).strip()[:120]
        if not title or not company:
            return jsonify({"ok": False, "error": "Missing title or company."}), 400
        job = {
            "job_id": linkedin_job_id(url), "title": title, "company": company,
            "location": str(data.get("location", "")).strip()[:120] or "India",
            "url": url, "source": "linkedin", "description": _fetch_description(url),
            "posted_at": str(data.get("posted_at", ""))[:40],
        }
    else:
        return jsonify({"ok": False, "error": "Unknown source."}), 400

    conn = get_conn(DB_PATH)
    try:
        dup = conn.execute("SELECT job_id FROM jobs WHERE user_id = ? AND (url = ? OR (LOWER(title) = ? AND LOWER(company) = ?))",
                           (uid, job["url"], job["title"].lower(), job["company"].lower())).fetchone()
        if dup:
            return jsonify({"ok": True, "already": True, "job_id": dup[0]})
        if not _insert_job(conn, job, user_id=uid):
            return jsonify({"ok": True, "already": True, "job_id": job["job_id"]})
        add_timeline(conn, job["job_id"], "Added from web search")
        conn.commit()
    finally:
        conn.close()
    run_in_background(("score", uid, job["job_id"]), _score_added_job, uid, job["job_id"])
    return jsonify({"ok": True, "job_id": job["job_id"]})


# ── Board Assistant (voice + chat) ─────────────────────────────────────────

@app.route("/api/voice/transcribe", methods=["POST"])
def voice_transcribe():
    uid = get_user_id()
    limited = _voice_rate_limited(uid, "transcribe", 20)
    if limited:
        return limited
    audio_file = request.files.get("file")
    if not audio_file:
        return jsonify({"ok": False, "error": "No audio received."}), 400
    audio_bytes = audio_file.read(ve.MAX_AUDIO_BYTES + 1)
    if not audio_bytes:
        return jsonify({"ok": False, "error": _t("I didn't catch any audio. Try again?", "मुझे कोई आवाज़ नहीं मिली। फिर से कोशिश करें?")}), 400
    if len(audio_bytes) > ve.MAX_AUDIO_BYTES:
        return jsonify({"ok": False, "error": _t("That recording was a bit long. Try a shorter question.", "रिकॉर्डिंग थोड़ी लंबी थी। छोटा सवाल पूछकर देखें।")}), 413
    mime_type = (audio_file.mimetype or "audio/webm").lower()
    g.voice_lang = ve.reply_lang(request.form.get("lang"))
    if mime_type not in ve.ALLOWED_AUDIO_TYPES:
        return jsonify({"ok": False, "error": "Unsupported audio format.", "provider_fallback": True}), 415

    conn = get_conn(DB_PATH)
    try:
        api_key = _voice_api_key(conn, uid)
    finally:
        conn.close()
    if not api_key:
        return jsonify({"ok": False, "error": "Voice transcription isn't set up.", "provider_fallback": True}), 400
    try:
        transcript = ve.transcribe_audio(audio_bytes, mime_type, api_key, lang=g.voice_lang)
    except ve.GeminiError as e:
        # The message never includes the API key; the key travels in a header
        app.logger.warning("Voice transcription failed for user %s: %s", uid, e)
        return jsonify({"ok": False, "error": _t("I couldn't transcribe that just now.", "अभी मैं इसे समझ नहीं पाई।"), "provider_fallback": True}), 502
    return jsonify({"ok": True, "transcript": transcript})


@app.route("/api/voice/synthesize", methods=["POST"])
def voice_synthesize():
    uid = get_user_id()
    limited = _voice_rate_limited(uid, "synthesize", 30)
    if limited:
        return limited
    d = request.get_json(silent=True) or {}
    text = str(d.get("text") or "").strip()
    if not text:
        return jsonify({"ok": False, "error": "No text provided"}), 400

    conn = get_conn(DB_PATH)
    try:
        api_key = _voice_api_key(conn, uid)
    finally:
        conn.close()
    if not api_key:
        return jsonify({"ok": False, "error": "Voice playback isn't set up.", "provider_fallback": True}), 400
    try:
        wav_bytes = ve.synthesize_speech(text, api_key)
    except ve.GeminiError as e:
        app.logger.warning("Voice synthesis failed for user %s: %s", uid, e)
        return jsonify({"ok": False, "error": "Voice playback is unavailable right now.", "provider_fallback": True}), 502
    return jsonify({"ok": True, "audio_base64": base64.b64encode(wav_bytes).decode("ascii")})


def shorten_title(title):
    if not title:
        return ""
    # Strip any suffix after dash, pipe, or parens
    for separator in (" - ", " — ", " | ", " ("):
        if separator in title:
            title = title.split(separator)[0]
    # Truncate if still too long
    if len(title) > 35:
        title = title[:32] + "..."
    return title.strip()

def clean_company(company):
    if not company:
        return ""
    if "demo_" in company:
        company = company.replace("demo_", "")
    for separator in ("_pm_", "_dev_", "_u"):
        if separator in company:
            company = company.split(separator)[0]
    if "_" in company:
        company = company.split("_")[0]
    
    special_cases = {
        "phonepe": "PhonePe",
        "razorpay": "Razorpay",
        "meesho": "Meesho",
        "groww": "Groww",
        "cred": "CRED",
        "swiggy": "Swiggy",
        "zepto": "Zepto",
        "itc": "ITC",
        "nike": "Nike"
    }
    company_lower = company.lower().strip()
    if company_lower in special_cases:
        return special_cases[company_lower]
    return company.capitalize()


@app.route("/api/voice/query", methods=["POST"])
def voice_query():
    uid = get_user_id()
    limited = _voice_rate_limited(uid, "query", 30)
    if limited:
        return limited
    d = request.get_json(silent=True) or {}
    transcript = ve.sanitize_transcript(d.get("transcript"))
    g.voice_lang = ve.reply_lang(d.get("lang"), transcript)
    if not transcript:
        return jsonify(_voice_reply(None, _t("I didn't hear anything that time. Tap the mic and try again?",
                                             "इस बार मुझे कुछ सुनाई नहीं दिया। माइक दबाकर फिर से बोलिए?"),
                                    suggestions=ve.follow_up_suggestions(None, lang=g.voice_lang)))
    history = ve.sanitize_history(d.get("chat_history"))

    conn = get_conn(DB_PATH)
    try:
        settings = get_user_settings(conn, uid)
        snapshot = _voice_jobs_snapshot(conn, uid)
        context_ids = ve.sanitize_job_ids(d.get("context_job_ids"), snapshot)
        res = ve.classify_intent_and_slot(transcript, snapshot, chat_history=history,
                                          api_key=settings.get("gemini_api_key"), context_job_ids=context_ids)
        res = _merge_voice_followup(res, session.pop("voice_awaiting", None), transcript, snapshot)
        return jsonify(_answer_voice_intent(conn, uid, settings, res, snapshot))
    except Exception:
        app.logger.exception("Voice query failed for user %s", uid)
        return jsonify(_voice_reply(None, _t("Sorry, something went wrong on my side. Mind trying that again?",
                                             "माफ़ कीजिए, मेरी तरफ़ से कुछ गड़बड़ हो गई। फिर से कोशिश करें?"),
                                    ok=False)), 500
    finally:
        conn.close()


def _merge_voice_followup(res, awaiting, transcript, snapshot):
    """Finish a request the assistant asked a follow-up question about.
    "Move it to applied" → "Which job?" → "Swiggy" should still move the Swiggy job."""
    if not awaiting or time.time() - awaiting.get("at", 0) > ve.PENDING_ACTION_TTL:
        return res
    slots = res["slots"]
    if awaiting.get("needs") == "job" and slots.get("job_id") and res["intent"] in (None, "job_lookup"):
        merged = dict(awaiting.get("slots") or {})
        merged.update({k: v for k, v in slots.items() if v})
        return {"intent": awaiting["intent"], "slots": merged, "ambiguous": res.get("ambiguous", False)}
    if awaiting.get("needs") == "status":
        column = slots.get("status") or slots.get("column") or ve._column_in(transcript.lower())
        job_id = (awaiting.get("slots") or {}).get("job_id")
        known = {str(j["job_id"]) for j in snapshot}
        if column and job_id in known and res["intent"] in (None, "column_count", "move_job", "job_lookup"):
            return {"intent": "move_job", "slots": {**ve._empty_slots(), "job_id": job_id, "status": column},
                    "ambiguous": False}
    return res


def _answer_voice_intent(conn, uid, settings, res, snapshot):
    intent = res.get("intent")
    slots = res.get("slots") or {}
    job = _voice_job(conn, uid, slots.get("job_id")) if slots.get("job_id") else None

    if intent in ve.JOB_INTENTS and not job:
        session["voice_awaiting"] = {"intent": intent, "slots": slots, "needs": "job", "at": time.time()}
        names = [ve._spoken_company(clean_company(j['company'])) for j in snapshot[:3]]
        chips = [_t(f"The {c} job", f"{c} वाली") for c in names]
        return _voice_reply(intent, ve.which_job_reply(intent, g.voice_lang), suggestions=chips)

    if intent in ve.MUTATING_INTENTS:
        return _propose_voice_action(conn, uid, intent, slots, job)

    cards = [_voice_card(job)] if job else []
    lang = g.voice_lang
    follow = ve.follow_up_suggestions(intent, job, lang=lang)

    if intent == "help":
        return _voice_reply(intent, ve.HELP_TEXT[lang], suggestions=follow)
    if intent == "greeting":
        return _voice_reply(intent, ve.greeting_text(settings.get("name"), lang=lang), suggestions=follow)
    if intent == "thanks":
        return _voice_reply(intent, _t(ve.pick("Anytime!", "Happy to help!", "You got it. Good luck out there!"),
                                       ve.pick("कभी भी!", "मदद करके खुशी हुई!", "बिल्कुल। ऑल द बेस्ट!")),
                            suggestions=follow)

    if intent == "daily_digest":
        return _digest_reply(conn, uid, mark_read=True)

    if intent == "pipeline_stats":
        by_status = dict(conn.execute(
            "SELECT status, COUNT(*) FROM jobs WHERE user_id = ? AND status != 'archived' GROUP BY status", (uid,)
        ).fetchall())
        active = sum(by_status.values())
        avg = conn.execute(
            "SELECT AVG(ai_score) FROM jobs WHERE user_id = ? AND status != 'archived' AND ai_score IS NOT NULL", (uid,)
        ).fetchone()[0] or 0
        counts = {
            "shortlisted": by_status.get("shortlisted", 0), "applied": by_status.get("applied", 0),
            "interviewing": by_status.get("interviewing", 0), "offers": by_status.get("offer", 0),
        }
        if not active:
            text = _t("Your board is empty right now. Say \"refresh my listings\" and I'll go find some jobs for you.",
                      "आपका बोर्ड अभी खाली है। \"लिस्टिंग रिफ्रेश करो\" कहिए और मैं आपके लिए जॉब्स ढूंढ लाऊँगी।")
        else:
            text = _t(f"You've got {active} job{'s' if active != 1 else ''} on your board, with an average match of "
                      f"{avg:.1f} out of 10. {counts['applied']} applied and {counts['interviewing']} interviewing.",
                      f"आपके बोर्ड पर {active} जॉब्स हैं, औसत मैच 10 में से {avg:.1f}। "
                      f"{counts['applied']} में अप्लाई किया है और {counts['interviewing']} में इंटरव्यू चल रहा है।")
            if counts["offers"]:
                text += _t(f" And {counts['offers']} offer{'s' if counts['offers'] != 1 else ''}. Nice work!",
                           f" और {counts['offers']} ऑफ़र भी! बहुत बढ़िया!")
            elif counts["applied"] == 0:
                text += _t(" Want me to pull up your top matches so you can start applying?",
                           " क्या मैं आपके टॉप मैच दिखाऊँ ताकि आप अप्लाई करना शुरू कर सकें?")
        card = {"type": "stats", "total": active, "avg_score": round(avg, 1), **counts}
        return _voice_reply(intent, text, cards=[card], suggestions=follow)

    if intent == "column_count":
        column = slots.get("column")
        if not column:
            return _voice_reply(intent, _t("Which column? New, Shortlisted, Interviewing, Applied, Offer or Rejected?",
                                           "कौन सा कॉलम? New, Shortlisted, Interviewing, Applied, Offer या Rejected?"),
                                suggestions=[ve.chip("shortlisted", lang), ve.chip("applied_count", lang)])
        statuses = ("new", "scored", "ready") if column == "new" else (column,)
        marks = ",".join("?" * len(statuses))
        count = conn.execute(f"SELECT COUNT(*) FROM jobs WHERE user_id = ? AND status IN ({marks})",
                             (uid, *statuses)).fetchone()[0]
        label = ve.column_label(column)
        if count == 0:
            text = _t(f"Nothing in {label} yet.", f"{label} में अभी कुछ नहीं है।")
        else:
            text = _t(f"You have {count} job{'s' if count != 1 else ''} in {label}.", f"{label} में आपकी {count} जॉब्स हैं।")
        return _voice_reply(intent, text, suggestions=follow)

    if intent == "job_lookup":
        label = ve.column_label(job["status"])
        if job["score"] is not None:
            text = _t(f"Here's {job['title']} at {job['company']}. It scored {job['score']:.1f} out of 10, and it's in your {label} column.",
                      f"यह रही {job['company']} की {job['title']} जॉब। इसका स्कोर 10 में से {job['score']:.1f} है, और यह {label} कॉलम में है।")
        else:
            text = _t(f"Here's {job['title']} at {job['company']}. I haven't scored it yet, and it's in your {label} column.",
                      f"यह रही {job['company']} की {job['title']} जॉब। इसका स्कोर अभी नहीं बना है, और यह {label} कॉलम में है।")
        return _voice_reply(intent, text, cards=cards, suggestions=follow)

    if intent == "job_fit":
        if job["score"] is None:
            text = _t(f"I haven't scored {job['title']} at {job['company']} yet. Try refreshing and I'll take a look.",
                      f"{job['company']} की जॉब का स्कोर अभी नहीं बना है। रिफ्रेश करके देखिए।")
        else:
            text = _t(f"{job['title']} at {job['company']} scored {job['score']:.1f} out of 10.",
                      f"{job['company']} की {job['title']} जॉब का स्कोर 10 में से {job['score']:.1f} है।")
            if job.get("summary"):
                text += " " + job["summary"].strip()
        return _voice_reply(intent, text, cards=cards, suggestions=follow)

    if intent == "job_status":
        status = job["status"]
        if status == "applied":
            text = _t(f"Yes, you've applied to {job['company']}. Fingers crossed!",
                      f"हाँ, आपने {job['company']} में अप्लाई कर दिया है। ऑल द बेस्ट!")
        elif status in ("interviewing", "offer"):
            text = _t(f"{job['company']} is in your {ve.column_label(status)} column. Exciting!",
                      f"{job['company']} आपके {ve.column_label(status)} कॉलम में है। वाह!")
        else:
            text = _t(f"Not yet. {job['title']} at {job['company']} is in your {ve.column_label(status)} column.",
                      f"अभी नहीं। {job['company']} की {job['title']} जॉब {ve.column_label(status)} कॉलम में है।")
        return _voice_reply(intent, text, cards=cards, suggestions=follow)

    if intent == "quiz_mode":
        reply = _voice_reply(intent, _t(f"Let's practice for {job['title']} at {job['company']}. Opening the quiz now.",
                                        f"चलिए {job['company']} की {job['title']} जॉब के लिए अभ्यास करते हैं। क्विज़ खोल रही हूँ।"))
        reply["action"] = {"type": "open_quiz", "job_id": job["job_id"]}
        return reply

    if intent == "email_lookup":
        company = clean_company(job["raw_company"]) if job else (slots.get("company") or "").strip()
        if company:
            rows = conn.execute(
                "SELECT sender, subject, body FROM received_emails WHERE user_id = ? AND (sender LIKE ? OR subject LIKE ? OR body LIKE ?) "
                "ORDER BY received_at DESC LIMIT 3",
                (uid, f"%{company}%", f"%{company}%", f"%{company}%")
            ).fetchall()
            text = (_t(f"Here's the latest from {company}: \"{rows[0][1]}\".", f"{company} का सबसे नया ईमेल: \"{rows[0][1]}\"।") if rows
                    else _t(f"Nothing from {company} yet. I'll keep an eye out.", f"{company} से अभी कुछ नहीं आया। मैं नज़र रखूँगी।"))
        else:
            rows = conn.execute(
                "SELECT sender, subject, body FROM received_emails WHERE user_id = ? ORDER BY received_at DESC LIMIT 3", (uid,)
            ).fetchall()
            text = _t("Here are your latest emails.", "ये रहे आपके सबसे नए ईमेल।") if rows else _t("No synced emails yet. Connect your inbox in Settings and I'll watch for replies.", "अभी कोई ईमेल सिंक नहीं हुआ। Settings में अपना इनबॉक्स जोड़िए, मैं जवाबों पर नज़र रखूँगी।")
        email_cards = [{"type": "email", "sender": r[0], "subject": r[1], "body": (r[2] or "")[:2000]} for r in rows]
        return _voice_reply(intent, text, cards=cards + email_cards, suggestions=follow)

    if intent == "email_count":
        count = conn.execute("SELECT COUNT(*) FROM received_emails WHERE user_id = ?", (uid,)).fetchone()[0]
        text = (_t(f"You have {count} synced email{'s' if count != 1 else ''}.", f"आपके {count} ईमेल सिंक हुए हैं।") if count
                else _t("No synced emails yet. Connect your inbox in Settings and I'll watch for replies.", "अभी कोई ईमेल सिंक नहीं हुआ। Settings में अपना इनबॉक्स जोड़िए, मैं जवाबों पर नज़र रखूँगी।"))
        return _voice_reply(intent, text, suggestions=follow)

    if intent == "last_sync":
        row = conn.execute("SELECT last_scraped_at FROM users WHERE id = ?", (uid,)).fetchone()
        when = _spoken_time(row[0], lang) if row and row[0] else None
        text = (_t(f"I last refreshed your board {when}.", f"आपका बोर्ड आखिरी बार {when} रिफ्रेश हुआ था।") if when
                else _t("Your board hasn't been refreshed yet. Want me to do it now?", "आपका बोर्ड अभी तक रिफ्रेश नहीं हुआ है। अभी कर दूँ?"))
        return _voice_reply(intent, text, suggestions=follow)

    if intent == "top_matches":
        rows = conn.execute(
            "SELECT job_id, title, company, ai_score, status, location, ai_summary FROM jobs WHERE user_id = ? AND status != 'archived' "
            "AND ai_score IS NOT NULL ORDER BY ai_score DESC LIMIT 3", (uid,)
        ).fetchall()
        if not rows:
            return _voice_reply(intent, _t("I haven't scored any jobs for you yet. Say \"refresh my listings\" and I'll get started.",
                                           "मैंने अभी तक आपकी कोई जॉब स्कोर नहीं की है। \"लिस्टिंग रिफ्रेश करो\" कहिए, मैं शुरू करती हूँ।"),
                                suggestions=[ve.chip("refresh", lang)])
        jobs = [_voice_job_from_row(r) for r in rows]
        if lang == "hi":
            spoken = ", ".join(f"{j['company']} की {j['title']} ({j['score']:.1f})" for j in jobs)
            text = f"आपके टॉप {len(jobs)} मैच: {spoken}।"
        else:
            spoken = ", ".join(f"{j['title']} at {j['company']} ({j['score']:.1f})" for j in jobs)
            text = f"Your top {'match is' if len(jobs) == 1 else str(len(jobs)) + ' matches are'}: {spoken}."
        return _voice_reply(intent, text, cards=[_voice_card(j) for j in jobs], suggestions=follow)

    if intent == "cover_letter_status":
        has_letter = conn.execute("SELECT 1 FROM cover_letters WHERE job_id = ? AND user_id = ?",
                                  (job["job_id"], uid)).fetchone()
        if has_letter:
            text = _t(f"Yes, there's a cover letter ready for {job['company']}.", f"हाँ, {job['company']} के लिए कवर लेटर तैयार है।")
            follow = [ve.chip("send", lang), ve.chip("rewrite", lang)]
        else:
            text = _t(f"Not yet. Want me to write one for {job['company']}?", f"अभी नहीं। क्या मैं {job['company']} के लिए एक लिख दूँ?")
            follow = [ve.chip("write_letter", lang)]
        return _voice_reply(intent, text, cards=cards, suggestions=follow)

    if intent in ("thumbs_up", "thumbs_down"):
        liked = intent == "thumbs_up"
        conn.execute("UPDATE jobs SET feedback = ? WHERE job_id = ? AND user_id = ?", (1 if liked else -1, job["job_id"], uid))
        conn.commit()
        rescore_inbox_in_background(uid)
        if liked:
            text = _t(f"Nice! I've marked {job['company']} as a favourite and I'll look for more jobs like it.",
                      f"बढ़िया! मैंने {job['company']} को पसंदीदा में डाल दिया है, ऐसी और जॉब्स ढूंढूँगी।")
        else:
            text = _t(f"Got it, you're not into {job['company']}. I'll show you fewer jobs like that.",
                      f"ठीक है, {job['company']} आपको पसंद नहीं। ऐसी जॉब्स कम दिखाऊँगी।")
        return _voice_reply(intent, text, cards=cards, suggestions=follow, board_changed=True)

    return _voice_reply(None, ve.pick(*ve.UNKNOWN_REPLIES[lang]), suggestions=ve.follow_up_suggestions(None, lang=lang))


def _propose_voice_action(conn, uid, intent, slots, job):
    """Changes always need a yes first. The pending action lives in the signed session with a
    one-time token, so a stale Confirm button can't fire a different action."""
    cards = [_voice_card(job)] if job else []
    if intent == "move_job":
        status = slots.get("status")
        if not status:
            session["voice_awaiting"] = {"intent": intent, "slots": slots, "needs": "status", "at": time.time()}
            return _voice_reply(intent, _t(f"Sure! Which column should {job['company']} go to?",
                                           f"ज़रूर! {job['company']} को किस कॉलम में डालूँ?"), cards=cards,
                                suggestions=["Shortlisted", "Applied", "Interviewing", "Rejected"])
        if job["status"] == status or (status == "new" and job["status"] in ("scored", "ready")):
            return _voice_reply(intent, _t(f"{job['company']} is already in {ve.column_label(status)}.",
                                           f"{job['company']} पहले से {ve.column_label(status)} में है।"), cards=cards)
        text = _t(f"Move {job['title']} at {job['company']} to {ve.column_label(status)}?",
                  f"{job['company']} की {job['title']} जॉब को {ve.column_label(status)} में डाल दूँ?")
    elif intent == "trigger_refresh":
        text = _t("Want me to refresh your listings and look for new jobs?", "क्या मैं आपकी लिस्टिंग रिफ्रेश करके नई जॉब्स ढूंढूँ?")
    elif intent == "regenerate_cover_letter":
        text = _t(f"Write a fresh cover letter for {job['title']} at {job['company']}?",
                  f"{job['company']} की {job['title']} जॉब के लिए नया कवर लेटर लिख दूँ?")
    elif intent == "send_email":
        has_letter = conn.execute("SELECT 1 FROM cover_letters WHERE job_id = ? AND user_id = ?",
                                  (job["job_id"], uid)).fetchone()
        if not has_letter:
            return _voice_reply(intent, _t(f"There's no cover letter for {job['company']} yet. Want me to write one first?",
                                           f"{job['company']} के लिए अभी कवर लेटर नहीं है। पहले एक लिख दूँ?"),
                                cards=cards, suggestions=[ve.chip("write_letter", g.voice_lang)])
        text = _t(f"Send the application email for {job['title']} at {job['company']}?",
                  f"{job['company']} की {job['title']} जॉब का एप्लीकेशन ईमेल भेज दूँ?")
    else:  # archive_job
        text = _t(f"Archive {job['title']} at {job['company']}? It'll disappear from your board.",
                  f"{job['company']} की {job['title']} जॉब आर्काइव कर दूँ? यह आपके बोर्ड से हट जाएगी।")

    token = secrets.token_urlsafe(12)
    session["voice_pending_action"] = {"intent": intent, "slots": slots, "token": token, "at": time.time()}
    reply = _voice_reply(intent, text, cards=cards)
    reply["requires_confirmation"] = True
    reply["action"] = {"intent": intent, "token": token}
    return reply


@app.route("/api/voice/confirm", methods=["POST"])
def voice_confirm():
    uid = get_user_id()
    limited = _voice_rate_limited(uid, "confirm", 15)
    if limited:
        return limited
    d = request.get_json(silent=True) or {}
    g.voice_lang = ve.reply_lang(d.get("lang"))
    pending = session.get("voice_pending_action")
    if not pending or d.get("token") != pending.get("token"):
        return jsonify({"ok": False, "error": _t("That request has expired. Just ask me again.", "वह अनुरोध पुराना हो गया। बस फिर से पूछिए।")}), 409
    session.pop("voice_pending_action", None)
    if time.time() - pending.get("at", 0) > ve.PENDING_ACTION_TTL:
        return jsonify({"ok": False, "error": _t("That request has expired. Just ask me again.", "वह अनुरोध पुराना हो गया। बस फिर से पूछिए।")}), 409

    intent = pending.get("intent")
    job_id = (pending.get("slots") or {}).get("job_id")
    conn = get_conn(DB_PATH)
    try:
        job = _voice_job(conn, uid, job_id) if job_id else None
        if intent in ve.JOB_INTENTS and not job:
            return jsonify({"ok": False, "error": _t("I couldn't find that job on your board anymore.", "वह जॉब अब आपके बोर्ड पर नहीं मिली।")}), 404

        if intent == "move_job":
            status = ve.normalize_column((pending.get("slots") or {}).get("status"))
            if not status:
                return jsonify({"ok": False, "error": _t("I lost track of which column you wanted. Ask me again?", "मैं भूल गई कि कौन सा कॉलम चाहिए था। फिर से बताइए?")}), 400
            conn.execute("UPDATE jobs SET status = ? WHERE job_id = ? AND user_id = ?", (status, job_id, uid))
            add_timeline(conn, job_id, f"Pipeline → {status} (Voice)")
            conn.commit()
            rescore_inbox_in_background(uid)
            text = _t(ve.pick("Done!", "All set!", "You got it!") + f" {job['company']} is now in {ve.column_label(status)}.",
                      ve.pick("हो गया!", "बिल्कुल!", "कर दिया!") + f" {job['company']} अब {ve.column_label(status)} में है।")
            return jsonify({"ok": True, "reply_text": text, "reply_cards": [_voice_card({**job, "status": status})],
                            "board_changed": True})

        if intent == "trigger_refresh":
            started = run_in_background(("sync", uid), run_user_sync, uid)
            text = (_t("On it! I'm looking for new jobs in the background. Your board will update when I'm done.",
                       "शुरू कर दिया! मैं पीछे से नई जॉब्स ढूंढ रही हूँ। काम होते ही आपका बोर्ड अपडेट हो जाएगा।")
                    if started else _t("I'm already refreshing your board. Hang tight!", "बोर्ड पहले से रिफ्रेश हो रहा है। थोड़ा रुकिए!"))
            return jsonify({"ok": True, "reply_text": text, "reply_cards": []})

        if intent == "regenerate_cover_letter":
            j = dict(conn.execute("SELECT * FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone())
            contact = conn.execute("SELECT name, title FROM contacts WHERE job_id = ? AND user_id = ? LIMIT 1", (job_id, uid)).fetchone()
            cn = contact[0] if contact else "Hiring Team"
            ct = contact[1] if contact else "Recruiter"
            settings = get_user_settings(conn, uid)
            api_key = settings.get("gemini_api_key")
            score_data = score_job(j["title"], j["company"], j.get("description") or "", resume_text=settings.get("resume_text"),
                                   api_key=api_key, user_id=uid)
            letter = generate_cover_letter(j["title"], j["company"], j.get("description") or "", cn, ct,
                                           resume_text=settings.get("resume_text"), api_key=api_key,
                                           candidate_name=settings.get("name") or "")
            li_note = generate_linkedin_note(cn, ct, j["company"], j["title"], api_key=api_key)
            conn.execute("UPDATE jobs SET ai_score = ?, ai_summary = ?, key_reqs = ? WHERE job_id = ? AND user_id = ?",
                         (score_data["score"], score_data["fit_summary"], json.dumps(score_data.get("key_requirements", [])), job_id, uid))
            conn.execute("INSERT OR REPLACE INTO cover_letters (job_id, subject, body, linkedin_note, created_at, user_id) VALUES (?,?,?,?,?,?)",
                         (job_id, letter["subject"], letter["body"], li_note, datetime.now().isoformat(), uid))
            add_timeline(conn, job_id, "Regenerated cover letter (Voice)")
            conn.commit()
            text = _t(f"Your new cover letter for {job['company']} is ready. Open the job to read it.",
                      f"{job['company']} के लिए आपका नया कवर लेटर तैयार है। पढ़ने के लिए जॉब खोलिए।")
            return jsonify({"ok": True, "reply_text": text, "board_changed": True,
                            "reply_cards": [_voice_card({**job, "score": score_data["score"]})]})

        if intent == "send_email":
            j = dict(conn.execute("SELECT * FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone())
            cl = conn.execute("SELECT * FROM cover_letters WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
            if not cl:
                return jsonify({"ok": False, "error": _t("There's no cover letter to send yet.", "भेजने के लिए अभी कोई कवर लेटर नहीं है।")}), 400
            cl = dict(cl)
            c = conn.execute("SELECT * FROM contacts WHERE job_id = ? AND user_id = ? LIMIT 1", (job_id, uid)).fetchone()
            c = dict(c) if c else {}
            item = {
                **j, "score": j.get("ai_score") or 0,
                "fit_summary": j.get("ai_summary") or "",
                "key_requirements": json.loads(j.get("key_reqs") or "[]"),
                "cover_letter_subject": cl["subject"],
                "cover_letter_body": cl["body"],
                "contact_name": c.get("name") or "Hiring Team",
                "contact_title": c.get("title") or "",
                "linkedin_note": cl.get("linkedin_note") or "",
            }
            if not send_email_digest([item], recipient=recipient_for_user(uid)):
                return jsonify({"ok": False, "error": _t("I couldn't send that email. Check your email settings and try again.", "ईमेल नहीं भेज पाई। अपनी ईमेल सेटिंग्स देखकर फिर से कोशिश करें।")}), 502
            conn.execute("UPDATE jobs SET status = 'applied' WHERE job_id = ? AND user_id = ?", (job_id, uid))
            add_timeline(conn, job_id, "Email sent → applied (Voice)")
            conn.commit()
            text = _t(f"Sent! I moved {job['company']} to Applied. Good luck!", f"भेज दिया! {job['company']} को Applied में डाल दिया। ऑल द बेस्ट!")
            return jsonify({"ok": True, "reply_text": text, "board_changed": True,
                            "reply_cards": [_voice_card({**job, "status": "applied"})]})

        if intent == "archive_job":
            conn.execute("UPDATE jobs SET status = 'archived' WHERE job_id = ? AND user_id = ?", (job_id, uid))
            add_timeline(conn, job_id, "Pipeline → archived (Voice)")
            conn.commit()
            text = _t(f"Archived {job['title']} at {job['company']}. One less thing to think about.",
                      f"{job['company']} की {job['title']} जॉब आर्काइव कर दी।")
            return jsonify({"ok": True, "reply_text": text, "reply_cards": [], "board_changed": True})

        return jsonify({"ok": False, "error": _t("I'm not sure what to do with that request.", "मुझे समझ नहीं आया कि इसका क्या करूँ।")}), 400
    except Exception:
        app.logger.exception("Voice action %s failed for user %s", intent, uid)
        return jsonify({"ok": False, "error": _t("Something went wrong while doing that. Please try again.", "यह करते समय कुछ गड़बड़ हो गई। कृपया फिर से कोशिश करें।")}), 500
    finally:
        conn.close()


@app.route("/api/voice/cancel", methods=["POST"])
def voice_cancel():
    session.pop("voice_pending_action", None)
    session.pop("voice_awaiting", None)
    g.voice_lang = ve.reply_lang((request.get_json(silent=True) or {}).get("lang"))
    return jsonify({"ok": True, "reply_text": _t(ve.pick("No problem, cancelled.", "Okay, I'll leave it as it is."),
                                                 ve.pick("ठीक है, रद्द कर दिया।", "ठीक है, जैसा है वैसा ही रहने देती हूँ।"))})


@app.route("/api/voice/welcome", methods=["GET"])
def voice_welcome():
    """Everything the assistant panel needs when it opens: a greeting, starter chips from the
    user's own board, the unread digest count, and whether server-side speech is available."""
    uid = get_user_id()
    lang = ve.reply_lang(request.args.get("lang"))
    conn = get_conn(DB_PATH)
    try:
        settings = get_user_settings(conn, uid)
        snapshot = _voice_jobs_snapshot(conn, uid)
        digest_count = len(_digest_rows(conn, uid))
        has_key = bool(settings.get("gemini_api_key") or get_fallback_gemini_key())
    finally:
        conn.close()
    return jsonify({
        "ok": True,
        "greeting": ve.greeting_text(settings.get("name"), lang=lang),
        "suggestions": ve.default_suggestions(snapshot, digest_count, lang=lang),
        "lang": lang,
        "digest_count": digest_count,
        "has_jobs": bool(snapshot),
        "server_voice": has_key,
    })


@app.route("/api/voice/digest", methods=["GET"])
def voice_digest():
    uid = get_user_id()
    conn = get_conn(DB_PATH)
    try:
        g.voice_lang = ve.reply_lang(request.args.get("lang"))
        return jsonify(_digest_reply(conn, uid, mark_read=request.args.get("peek") != "true"))
    finally:
        conn.close()


# ── Board Assistant helpers ────────────────────────────────────────────────

def _voice_rate_limited(uid, bucket, limit):
    if ve.rate_limiter.allow(uid, bucket, limit):
        return None
    return jsonify({"ok": False, "rate_limited": True,
                    "error": _t("You're going a little fast for me. Give me a few seconds and try again.",
                                "आप थोड़ा तेज़ चल रहे हैं। कुछ सेकंड रुककर फिर से कोशिश करें।")}), 429


def _voice_api_key(conn, uid):
    return get_user_settings(conn, uid).get("gemini_api_key") or get_fallback_gemini_key()


def _voice_reply(intent, text, cards=None, suggestions=None, ok=True, **extra):
    cards = cards or []
    return {
        "ok": ok,
        "intent": intent,
        "reply_text": text,
        "requires_confirmation": False,
        "reply_cards": cards,
        "suggestions": suggestions or [],
        "context_job_ids": [c["job_id"] for c in cards if c.get("type") == "job"],
        "lang": getattr(g, "voice_lang", "en"),
        **extra,
    }


def _voice_jobs_snapshot(conn, uid):
    rows = conn.execute(
        "SELECT job_id, title, company, status, ai_score FROM jobs WHERE user_id = ? AND status != 'archived' ORDER BY scraped_at DESC LIMIT 50",
        (uid,)
    ).fetchall()
    return [{"job_id": str(r[0]), "title": r[1] or "", "company": r[2] or "", "status": r[3], "score": r[4]} for r in rows]


def _voice_job_from_row(r):
    return {"job_id": str(r[0]), "title": shorten_title(r[1]), "company": clean_company(r[2]), "raw_company": r[2] or "",
            "score": r[3], "status": r[4], "location": r[5] or "", "summary": r[6] or ""}


def _voice_job(conn, uid, job_id):
    row = conn.execute(
        "SELECT job_id, title, company, ai_score, status, location, ai_summary FROM jobs WHERE job_id = ? AND user_id = ?",
        (job_id, uid)
    ).fetchone()
    return _voice_job_from_row(row) if row else None


def _voice_card(job):
    return {"type": "job", "job_id": job["job_id"], "title": job["title"], "company": job["company"],
            "score": job["score"], "status": job["status"], "location": job["location"]}


def _spoken_time(value, lang="en"):
    try:
        dt = datetime.fromisoformat(str(value).split(".")[0])
    except ValueError:
        return None
    days = (datetime.now().date() - dt.date()).days
    clock = dt.strftime("%I:%M %p").lstrip("0")
    if lang == "hi":
        day = "आज" if days == 0 else "कल" if days == 1 else dt.strftime("%d %b")
        return f"{day} {clock} बजे"
    if days == 0:
        return f"today at {clock}"
    if days == 1:
        return f"yesterday at {clock}"
    return f"on {dt.strftime('%b %d')} at {clock}"


def _t(en, hi):
    """Pick the assistant's wording for this request's language (set from the client and the transcript)."""
    return hi if getattr(g, "voice_lang", "en") == "hi" else en


def _digest_rows(conn, uid):
    row = conn.execute("SELECT last_digest_read_at FROM users WHERE id = ?", (uid,)).fetchone()
    since = row[0] if row and row[0] else (datetime.now() - timedelta(days=1)).isoformat()
    return conn.execute(
        "SELECT job_id, title, company, ai_score, status, location, ai_summary FROM jobs WHERE user_id = ? AND scraped_at > ? "
        "AND status != 'archived' ORDER BY ai_score DESC LIMIT 5",
        (uid, since)
    ).fetchall()


def _digest_reply(conn, uid, mark_read):
    jobs = [_voice_job_from_row(r) for r in _digest_rows(conn, uid)]
    if mark_read:
        conn.execute("UPDATE users SET last_digest_read_at = ? WHERE id = ?", (datetime.now().isoformat(), uid))
        conn.commit()
    lang = getattr(g, "voice_lang", "en")
    if not jobs:
        text = _t("You're all caught up! No new recommendations since you last checked.",
                  "सब देख लिया! पिछली बार के बाद कोई नई सिफारिश नहीं है।")
        follow = [ve.chip("top", lang), ve.chip("refresh", lang)]
    else:
        top = jobs[0]
        if lang == "hi":
            score = f", स्कोर {top['score']:.1f}" if top["score"] is not None else ""
            text = f"खुशखबरी! आपके लिए {len(jobs)} नई सिफारिशें हैं। सबसे अच्छी है {top['company']} की {top['title']}{score}।"
            if len(jobs) > 1:
                text += " और नई: " + ", ".join(f"{j['company']} की {j['title']}" for j in jobs[1:]) + "।"
        else:
            score = f" with a score of {top['score']:.1f}" if top["score"] is not None else ""
            text = (f"Good news! You have {len(jobs)} new recommendation{'s' if len(jobs) != 1 else ''}. "
                    f"The best one is {top['title']} at {top['company']}{score}.")
            if len(jobs) > 1:
                text += " Also new: " + ", ".join(f"{j['title']} at {j['company']}" for j in jobs[1:]) + "."
        follow = [ve.chip("first", lang), ve.chip("first_why", lang)]
    reply = _voice_reply("daily_digest", text, cards=[_voice_card(j) for j in jobs], suggestions=follow)
    reply.update({"count": len(jobs), "jobs": [{k: j[k] for k in ("job_id", "title", "company", "score")} for j in jobs]})
    return reply


def scrape_job_url(url: str) -> dict:
    """Fetches a job URL and extracts company, title, description, and location."""
    import requests
    from bs4 import BeautifulSoup
    import re
    import json
    from scraper import HEADERS, _fetch_description
    from ai_engine import GEMINI_API_KEY, _call_gemini
    
    domain = url.split("://")[-1].split("/")[0].replace("www.", "")
    company = domain.split(".")[0].capitalize()
    title = "Product Manager"
    desc = "Forwarded via WhatsApp."
    loc = "India"
    
    try:
        resp = requests.get(url, headers=HEADERS, timeout=10)
        if resp.status_code == 200:
            soup = BeautifulSoup(resp.text, "html.parser")
            fetched_desc = _fetch_description(url)
            if fetched_desc:
                desc = fetched_desc
            else:
                desc = soup.get_text(separator="\n", strip=True)[:2500]
                
            if GEMINI_API_KEY:
                page_text = soup.get_text(separator="\n", strip=True)[:4000]
                prompt = f"""
Analyze the following text content fetched from a job posting webpage ({url}).
Extract the:
- Company name (e.g. "Adobe", "Razorpay")
- Job Title (e.g. "Senior Product Manager", "Product Manager - Payments")
- Job Location (e.g. "Bengaluru", "Remote", "India")
- Brief job description (max 1000 characters)

Page Text:
{page_text}

Respond ONLY with a JSON object containing keys: "company", "title", "location", "description". No markdown wrapping or extra comments.
"""
                try:
                    res = _call_gemini(prompt)
                    clean_res = res.replace("```json", "").replace("```", "").strip()
                    parsed_data = json.loads(clean_res)
                    if parsed_data.get("company"):
                        company = parsed_data["company"].strip()
                    if parsed_data.get("title"):
                        title = parsed_data["title"].strip()
                    if parsed_data.get("location"):
                        loc = parsed_data["location"].strip()
                    if parsed_data.get("description"):
                        desc = parsed_data["description"].strip()
                except Exception as e:
                    print(f"[scrape_job_url] Gemini extraction failed: {e}")
            
            if not GEMINI_API_KEY or company == domain.split(".")[0].capitalize():
                # Fallback meta extraction
                og_title = soup.find("meta", property="og:title")
                title_text = og_title["content"] if og_title and og_title.get("content") else (soup.title.string if soup.title else "")
                if title_text:
                    title_text = title_text.strip()
                    if " hiring " in title_text:
                        parts = title_text.split(" hiring ")
                        company = parts[0].strip()
                        title = parts[1].split(" in ")[0].split("|")[0].split("-")[0].strip()
                    elif " at " in title_text:
                        parts = title_text.split(" at ")
                        title = parts[0].strip()
                        company = parts[1].split("|")[0].split("-")[0].strip()
                    elif " - " in title_text:
                        parts = title_text.split(" - ")
                        company = parts[0].strip()
                        title = parts[1].split("|")[0].strip()
                    else:
                        title = title_text.split("|")[0].split("-")[0].strip()
    except Exception as e:
        print(f"[scrape_job_url] Error: {e}")
        
    return {
        "company": company,
        "title": title,
        "url": url,
        "location": loc,
        "description": desc
    }


@app.route("/api/whatsapp/import", methods=["POST"])
def import_whatsapp():
    text = request.json.get("text", "")
    if not text:
        return jsonify({"ok": False, "error": "Empty text"}), 400
        
    prompt = f"""
Given the following raw chat message(s) shared on WhatsApp, extract all job listings mentioned.
For each job listing, extract:
- Company name (clean, standard name, e.g. "Swiggy" instead of "Swiggy link")
- Job title (clean, standard title, e.g. "Product Manager", "Software Engineer")
- Job URL / link (if present, else empty string)
- Brief description / context from the message (if present, else empty string)

Raw chat text:
{text}

Respond ONLY with a JSON array of objects with the keys: "company", "title", "url", "description". Do not include markdown wraps or any explanation.
"""
    import json
    parsed = []
    try:
        from ai_engine import GEMINI_API_KEY, _call_gemini
        if GEMINI_API_KEY:
            raw_response = _call_gemini(prompt)
            clean_response = raw_response.replace("```json", "").replace("```", "").strip()
            parsed = json.loads(clean_response)
    except Exception as e:
        print(f"[WhatsApp Import] Gemini failed, falling back to regex: {e}")
        
    # Enrich by scraping URLs if present
    enriched_parsed = []
    for item in parsed:
        url = item.get("url", "").strip()
        if url:
            print(f"[WhatsApp Import] Enrichment scraping URL: {url}")
            scraped = scrape_job_url(url)
            item["company"] = scraped["company"]
            item["title"] = scraped["title"]
            item["description"] = scraped["description"]
            item["location"] = scraped.get("location", "Remote")
        enriched_parsed.append(item)
    parsed = enriched_parsed

    # Regex fallback if Gemini failed or is not active
    if not parsed:
        import re
        urls = re.findall(r'https?://[^\s<>"]+|www\.[^\s<>"]+', text)
        for url in urls:
            print(f"[WhatsApp Import] Scraping fallback URL: {url}")
            scraped = scrape_job_url(url)
            parsed.append(scraped)
            
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    import uuid
    import random
    from ai_engine import score_job
    
    added_count = 0
    for item in parsed:
        company = item.get("company", "Unknown Company").strip()
        title = item.get("title", "Product Manager").strip()
        url = item.get("url", "").strip()
        desc = item.get("description", "Imported from WhatsApp.").strip()
        loc = item.get("location", "Remote").strip()
        
        if url:
            existing = conn.execute(
                "SELECT title, company, status FROM jobs WHERE url = ? AND user_id = ?", (url, uid)
            ).fetchone()
        if not existing:
            existing = conn.execute(
                "SELECT title, company, status FROM jobs WHERE LOWER(company) = ? AND LOWER(title) = ? AND user_id = ?",
                (company.lower(), title.lower(), uid)
            ).fetchone()
            
        if existing:
            status_map = {
                "whatsapp": "From WhatsApp",
                "new": "Inbox / New", "scored": "Inbox / New", "ready": "Inbox / New",
                "shortlisted": "Shortlisted / Active", "interviewing": "Shortlisted / Active",
                "applied": "Applied", "offer": "Offer", "rejected": "Rejected", "archived": "Archived"
            }
            lane = status_map.get(existing["status"], "Unknown")
            conn.close()
            return jsonify({
                "ok": False,
                "error": f"Duplicate found: '{existing['title']}' at {existing['company']} is already in the '{lane}' lane!"
            }), 409

        job_id = f"wa_{company.lower().replace(' ', '_')}_{str(uuid.uuid4())[:8]}"
        
        # Calculate real score if description and title are present
        try:
            settings = get_user_settings(conn, uid)
            score_data = score_job(title, company, desc, resume_text=settings.get("resume_text"), api_key=settings.get("gemini_api_key"), user_id=uid)
            auto_score = score_data.get("score", 7.0)
            ai_summary = score_data.get("summary", "Imported from WhatsApp.")
            key_reqs = json.dumps(score_data.get("key_requirements", []))
        except Exception as e:
            print(f"[WhatsApp Import] Scoring failed: {e}")
            auto_score = round(random.uniform(7.5, 9.5), 1)
            ai_summary = desc
            key_reqs = "[]"
            
        conn.execute(
            """INSERT INTO jobs (job_id, company, title, status, url, location, description, scraped_at, ai_score, ai_summary, key_reqs, user_id)
               VALUES (?, ?, ?, 'whatsapp', ?, ?, ?, ?, ?, ?, ?, ?)""",
            (job_id, company, title, url, loc, desc, datetime.now().isoformat(), auto_score, ai_summary, key_reqs, uid)
        )
        add_timeline(conn, job_id, "Imported from WhatsApp forward")
        added_count += 1
            
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "count": added_count})


@app.route("/api/job/<job_id>/resume", methods=["GET"])
def get_resume(job_id):
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    row = conn.execute("SELECT resume_content FROM tailored_resumes WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
    
    if row:
        resume = row["resume_content"]
    else:
        # Fetch job details to generate
        job = conn.execute("SELECT title, company, description FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
        if not job:
            conn.close()
            return jsonify({"ok": False, "error": "Job not found"}), 404
            
        from ai_engine import generate_tailored_resume
        settings = get_user_settings(conn, uid)
        resume = generate_tailored_resume(job["description"] or "", job["title"] or "", job["company"] or "", resume_text=settings.get("resume_text"), api_key=settings.get("gemini_api_key"))
        
        # Save to DB cache
        conn.execute(
            "INSERT OR REPLACE INTO tailored_resumes (job_id, resume_content, created_at, user_id) VALUES (?, ?, ?, ?)",
            (job_id, resume, datetime.now().isoformat(), uid)
        )
        conn.commit()
        
    conn.close()
    return jsonify({"ok": True, "resume": resume})


@app.route("/resume/<job_id>/print")
def print_resume(job_id):
    conn = get_conn(DB_PATH)
    uid = get_user_id()
    row = conn.execute("SELECT resume_content FROM tailored_resumes WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
    
    if not row:
        job = conn.execute("SELECT title, company, description FROM jobs WHERE job_id = ? AND user_id = ?", (job_id, uid)).fetchone()
        if not job:
            conn.close()
            return "Job not found", 404
        from ai_engine import generate_tailored_resume
        settings = get_user_settings(conn, uid)
        resume = generate_tailored_resume(job["description"] or "", job["title"] or "", job["company"] or "", resume_text=settings.get("resume_text"), api_key=settings.get("gemini_api_key"))
        conn.execute(
            "INSERT OR REPLACE INTO tailored_resumes (job_id, resume_content, created_at, user_id) VALUES (?, ?, ?, ?)",
            (job_id, resume, datetime.now().isoformat(), uid)
        )
        conn.commit()
    else:
        resume = row["resume_content"]
    conn.close()
    
    # Preprocess markdown to ensure valid list blocks and header spacing
    lines = resume.splitlines()
    new_lines = []
    for i, line in enumerate(lines):
        clean = line.strip()
        if clean.startswith(("-", "*")):
            if i > 0 and lines[i-1].strip() and not lines[i-1].strip().startswith(("-", "*")):
                new_lines.append("")
        elif clean.startswith("**") and "float: right" in line:
            if i > 0 and lines[i-1].strip():
                new_lines.append("")
        new_lines.append(line)
    processed_resume = "\n".join(new_lines)
    
    import markdown
    html_content = markdown.markdown(processed_resume)
    
    return f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <title>Tailored Resume - {job_id}</title>
  <style>
    @page {{
      size: letter;
      margin: 0.6in 0.6in 0.5in 0.6in;
    }}
    body {{
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      color: #000;
      line-height: 1.45;
      max-width: 800px;
      margin: 0 auto;
      background: #fff;
      font-size: 11px;
    }}
    h1 {{
      font-size: 22px;
      text-align: center;
      margin-top: 0;
      margin-bottom: 2px;
      font-weight: bold;
      color: #000;
      text-transform: none;
      letter-spacing: -0.2px;
    }}
    /* Style the contact information paragraph under the name */
    .resume-body > p:first-of-type {{
      text-align: center;
      font-size: 10px;
      margin-top: 0;
      margin-bottom: 20px;
      color: #333;
      word-spacing: 1px;
    }}
    h2 {{
      border-bottom: 1px solid #111;
      padding-bottom: 2px;
      margin-top: 20px;
      margin-bottom: 10px;
      font-size: 11px;
      text-transform: uppercase;
      letter-spacing: 0.8px;
      font-weight: bold;
      color: #000;
    }}
    p {{
      margin-top: 12px;
      margin-bottom: 4px;
      font-size: 11px;
    }}
    /* Tighten gap between title paragraph and company/context paragraph */
    .resume-body p + p {{
      margin-top: 2px;
      margin-bottom: 4px;
    }}
    ul {{
      padding-left: 18px;
      margin-top: 4px;
      margin-bottom: 12px;
    }}
    li {{
      margin-top: 2px;
      margin-bottom: 3px;
      font-size: 11px;
      line-height: 1.4;
    }}
    .print-btn-container {{
      text-align: center;
      margin-bottom: 20px;
      margin-top: 10px;
    }}
    .btn {{
      padding: 8px 16px;
      background: #5c3bcf;
      color: white;
      border: none;
      border-radius: 4px;
      cursor: pointer;
      font-size: 12px;
      font-family: sans-serif;
      font-weight: 500;
      box-shadow: 0 1px 3px rgba(0,0,0,0.15);
    }}
    @media print {{
      body {{
        padding: 0;
        margin: 0;
      }}
      .print-btn-container {{
        display: none !important;
      }}
    }}
  </style>
</head>
<body>
  <div class="print-btn-container">
    <button onclick="window.print()" class="btn">🖨️ Save / Download as PDF</button>
  </div>
  <div class="resume-body">
    {html_content}
  </div>
</body>
</html>"""


@app.route("/logo.svg")
def serve_logo():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return send_file(os.path.join(base_dir, "logo.svg"), mimetype="image/svg+xml")


@app.route("/favicon.ico")
def serve_favicon():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return send_file(os.path.join(base_dir, "logo.svg"), mimetype="image/svg+xml")


@app.route("/hero_background.png")
def serve_hero_background():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return send_file(os.path.join(base_dir, "hero_background.png"), mimetype="image/png")


@app.route("/api/debug/db")
def debug_db():
    if not is_admin():
        return "Unauthorized", 401
    from db import IS_POSTGRES, DATABASE_URL
    import re
    conn = get_conn(DB_PATH)
    try:
        user_count = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        job_count = conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
        jobs_sample = conn.execute("SELECT job_id, title, status, user_id FROM jobs LIMIT 5").fetchall()
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)})
    finally:
        conn.close()
        
    db_masked = ""
    if IS_POSTGRES:
        db_masked = re.sub(r":([^@]+)@", ":***@", DATABASE_URL)
        
    return jsonify({
        "ok": True,
        "is_postgres": IS_POSTGRES,
        "database_url_masked": db_masked,
        "db_path": DB_PATH,
        "user_count": user_count,
        "job_count": job_count,
        "jobs_sample": [dict(r) for r in jobs_sample]
    })


def run_dashboard(port=5050):
    init_db(DB_PATH)
    print(f"\nDashboard → http://localhost:{port}")
    print("Press Ctrl+C to stop.\n")
    app.run(debug=False, port=port, use_reloader=False)


if __name__ == "__main__":
    run_dashboard()
